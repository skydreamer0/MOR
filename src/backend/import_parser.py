"""Bounded, exact-schema workbook discovery and lossless typed validation.

No filename inference, fuzzy headers, date serial guesses, or silent coercion.
Coordinates refer to the supplied workbook, including moved table headers.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
import hashlib
import math
import numbers
import re
import zipfile

import openpyxl
import pandas as pd

SALES_HEADERS = {'年','月','日','客戶簡稱','商品號','商品簡稱','銷+贈S量','單價NT(淨)','含稅總額(淨)'}
SHPB_HEADERS = {'出貨日期','客戶簡稱','產品','產品簡稱','銷售數量','贈品數量','銷貨淨價','含稅淨額'}
MONTHLY_HEADERS = {'客戶簡稱','商品號','標  題'}
BUDGET_LABELS = {'02預算總量','04預算總額','07預算S總量'}
HISTORY_LABELS = {'01銷+贈量','02銷+贈S量','03銷含稅總額','04銷未稅總額','05折後業績額'}
MAX_BYTES = 25 * 1024 * 1024
MAX_UNPACKED_BYTES = 200 * 1024 * 1024
MAX_CELLS = 500_000
MAX_SHEETS = 50


def is_blank(value):
    return value is None or (isinstance(value, str) and not value.strip()) or (not isinstance(value, str) and bool(pd.isna(value)))


def strict_int(value, label='integer'):
    if isinstance(value, bool) or is_blank(value):
        raise ValueError(f'{label}: required integral value')
    if isinstance(value, numbers.Real):
        if not math.isfinite(float(value)) or float(value) != int(value):
            raise ValueError(f'{label}: fractional/nonfinite value {value!r}')
        return int(value)
    if isinstance(value, str) and re.fullmatch(r'[0-9]+', value):
        return int(value)
    raise ValueError(f'{label}: invalid integer {value!r}')


def strict_number(value, label='number'):
    if isinstance(value, bool) or is_blank(value):
        raise ValueError(f'{label}: blank is not zero')
    if isinstance(value, str) and not re.fullmatch(r'[+-]?(?:\d+(?:\.\d*)?|\.\d+)', value):
        raise ValueError(f'{label}: invalid numeric value {value!r}')
    try:
        result = float(value)
    except (TypeError, ValueError):
        raise ValueError(f'{label}: invalid numeric value {value!r}') from None
    if not math.isfinite(result):
        raise ValueError(f'{label}: number must be finite')
    return result


def strict_identifier(value, label='identifier'):
    if is_blank(value) or isinstance(value, bool):
        raise ValueError(f'{label}: missing identifier')
    if isinstance(value, str):
        if value != value.strip():
            raise ValueError(f'{label}: surrounding whitespace requires correction')
        return value
    return str(strict_int(value, label))


def strict_date(value, label='date'):
    if is_blank(value):
        raise ValueError(f'{label}: missing date')
    if isinstance(value, (datetime, pd.Timestamp)):
        if value.tzinfo is not None or value.time() != datetime.min.time():
            raise ValueError(f'{label}: expected an unzoned calendar date')
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, numbers.Real) and not isinstance(value, bool):
        value = str(strict_int(value, label))
    if isinstance(value, str):
        for pattern, fmt in ((r'\d{8}', '%Y%m%d'), (r'\d{4}-\d{2}-\d{2}', '%Y-%m-%d')):
            if re.fullmatch(pattern, value):
                try:
                    return datetime.strptime(value, fmt).date()
                except ValueError:
                    break
    raise ValueError(f'{label}: invalid date {value!r}; use YYYYMMDD, YYYY-MM-DD or Excel date cells')


def parse_period(header, row_year=None, default_year=None):
    period = strict_int(header, 'month header')
    encoded_year = None
    if 100001 <= period <= 999912:
        encoded_year, month = divmod(period, 100)
    else:
        month = period
    if not 1 <= month <= 12:
        raise ValueError(f'Invalid month header {header!r}')
    years = [strict_int(v, 'year') for v in (encoded_year,row_year,default_year) if v is not None and not is_blank(v)]
    if not years or len(set(years)) != 1 or not 1900 <= years[0] <= 9999:
        raise ValueError(f'Missing/conflicting year for month header {header!r}')
    return years[0], month


@dataclass(frozen=True)
class ParsedTable:
    kind: str
    frame: pd.DataFrame
    source_filename: str
    source_hash: str
    sheet: str
    header_row: int
    header_column: int

    def source(self):
        return {'filename':self.source_filename,'sha256':self.source_hash,'sheet':self.sheet,'header_row':self.header_row,'header_column':self.header_column}


def _signature(headers):
    names = set(headers)
    kinds = [kind for kind, required in [('sales', SALES_HEADERS),('shpb',SHPB_HEADERS),('monthly',MONTHLY_HEADERS)] if required <= names]
    if len(kinds) > 1:
        raise ValueError('Ambiguous table contains multiple incompatible schemas')
    return kinds[0] if kinds else None


def parse_workbook(payload: bytes, source_filename: str) -> list[ParsedTable]:
    if len(payload) > MAX_BYTES:
        raise ValueError(f'{source_filename}: workbook exceeds {MAX_BYTES} bytes')
    try:
        with zipfile.ZipFile(BytesIO(payload)) as archive:
            if sum(x.file_size for x in archive.infolist()) > MAX_UNPACKED_BYTES:
                raise ValueError('Expanded workbook exceeds size limit')
        book = openpyxl.load_workbook(BytesIO(payload), read_only=True, data_only=False)
        cached = openpyxl.load_workbook(BytesIO(payload), read_only=True, data_only=True)
    except (zipfile.BadZipFile, OSError) as exc:
        raise ValueError(f'{source_filename}: expected an xlsx workbook') from exc
    try:
        if len(book.worksheets) > MAX_SHEETS:
            raise ValueError('Workbook exceeds sheet limit')
        if sum((s.max_row or 0)*(s.max_column or 0) for s in book) > MAX_CELLS:
            raise ValueError('Workbook exceeds scanned-cell limit')
        tables = []
        scanned_cells = 0
        digest = hashlib.sha256(payload).hexdigest()
        for sheet in book:
            # Do not trust understated worksheet dimensions to truncate facts.
            sheet.reset_dimensions()
            grid = []
            width = 0
            for row in sheet.iter_rows():
                width = max(width, len(row))
                if scanned_cells + (len(grid) + 1) * width > MAX_CELLS:
                    raise ValueError('Workbook exceeds scanned-cell limit')
                grid.append(list(row))
            scanned_cells += len(grid) * width
            grid = [row + [None] * (width-len(row)) for row in grid]
            values = [[c.value if c is not None else None for c in row] for row in grid]
            candidates = []
            for ri, row in enumerate(values):
                # Contiguous header regions allow side-by-side, non-overlapping tables.
                starts = [i for i,v in enumerate(row) if v is not None and (i == 0 or row[i-1] is None)]
                for start in starts:
                    end = start
                    while end < len(row) and row[end] is not None:
                        end += 1
                    headers = row[start:end]
                    kind = _signature(headers)
                    if not kind and '客戶簡稱' in headers and ({'商品號','產品'} & set(headers)):
                        raise ValueError(f'{source_filename}/{sheet.title}!{ri+1}: incomplete supported schema; missing required headers')
                    if kind:
                        if len(set(headers)) != len(headers):
                            raise ValueError(f'{source_filename}/{sheet.title}!{ri+1}: duplicate headers')
                        candidates.append((ri,start,end,headers,kind))
            cached_rows = None
            for ri,start,end,headers,kind in candidates:
                stop = min([other[0] for other in candidates if other[0]>ri and max(start,other[1])<min(end,other[2])] or [len(values)])
                data, row_numbers = [], []
                for r in range(ri+1,stop):
                    raw = values[r][start:end]
                    if not any(v is not None for v in raw):
                        continue
                    # Exact documented summary row rule; other malformed rows fail later.
                    if kind == 'shpb' and raw[headers.index('產品簡稱')] == '每日總計' and all(is_blank(raw[headers.index(k)]) for k in ('出貨日期','客戶簡稱','產品')):
                        continue
                    for offset, val in enumerate(raw):
                        cell = grid[r][start+offset]
                        if cell is None:
                            continue
                        if headers[offset] in {'商品號','產品','客戶代號'} and isinstance(val,numbers.Real) and re.fullmatch(r'0{2,}',cell.number_format or ''):
                            raise ValueError(f'{source_filename}/{sheet.title}!{cell.coordinate}: formatted numeric identifier; supply text to preserve leading zeros')
                        if cell.data_type == 'e':
                            raise ValueError(f'{source_filename}/{sheet.title}!{cell.coordinate}: formula error')
                        if cell.data_type == 'f':
                            if cached_rows is None:
                                cached[sheet.title].reset_dimensions()
                                cached_rows = list(cached[sheet.title].iter_rows(values_only=True))
                            cached_row = cached_rows[r] if r < len(cached_rows) else ()
                            val = cached_row[start+offset] if start+offset < len(cached_row) else None
                            if val is None:
                                raise ValueError(f'{source_filename}/{sheet.title}!{cell.coordinate}: formula has no cached result')
                            raw[offset] = val
                    data.append(raw)
                    row_numbers.append(r+1)
                if not data:
                    raise ValueError(f'{source_filename}/{sheet.title}!{ri+1}: empty recognized table')
                frame = pd.DataFrame(data, columns=headers, dtype=object)
                frame.attrs['source'] = f'{source_filename}/{sheet.title}'
                frame.attrs['rows'] = row_numbers
                frame.attrs['columns'] = {h:openpyxl.utils.get_column_letter(start+i+1) for i,h in enumerate(headers)}
                if kind == 'monthly':
                    labels=set(frame['標  題'])
                    budget=bool(labels & BUDGET_LABELS); history=bool(labels & (HISTORY_LABELS - {'01銷+贈量'})) or (not budget and bool(labels & HISTORY_LABELS))
                    if budget and history:
                        raise ValueError(f'{source_filename}/{sheet.title}: mixed budget/history metrics require separate tables')
                    if not budget and not history:
                        raise ValueError(f'{source_filename}/{sheet.title}: unknown monthly metric labels')
                    kind = 'budget' if budget else 'monthly_actual'
                tables.append(ParsedTable(kind,frame,source_filename,digest,sheet.title,ri+1,start+1))
                if len(tables)>100:
                    raise ValueError('Workbook exceeds 100 recognized tables')
        if not tables:
            raise ValueError(f'{source_filename}: no recognized tables; transpose and fuzzy mapping are unsupported')
        return tables
    finally:
        book.close();cached.close()


def cell_label(frame, index, column):
    rows=frame.attrs.get('rows', [])
    row=rows[index] if index < len(rows) else index+2
    col=frame.attrs.get('columns', {}).get(column,str(column))
    return f"{frame.attrs.get('source','table')}!{col}{row} ({column})"
