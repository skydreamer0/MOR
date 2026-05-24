from datetime import date
from io import BytesIO

import openpyxl

from src.backend.exporter import export_forecast
from src.backend.forecast_models import ForecastRow, ForecastSummary


def _make_row(*, excluded: bool = False) -> ForecastRow:
    return ForecastRow(
        row_id="A__P1",
        customer="A醫院",
        product_code="P1",
        product_name="產品一",
        latest_order_date=date(2026, 4, 1),
        cycle_days=30,
        next_order_date=date(2026, 5, 1),
        auto_in_month=True,
        last_year_same_month_qty=10.0,
        this_year_same_month_qty=12.0,
        latest_price=100.0,
        system_forecast=12.0,
        manual_adjustment=None,
        final_forecast=12.0,
        estimated_amount=1200.0,
        forecast_basis="last_year",
        budget_quantity=10.0,
        excluded=excluded,
    )


def _make_summary(rows=None) -> ForecastSummary:
    if rows is None:
        rows = [_make_row()]
    return ForecastSummary(year=2026, month=5, rows=rows, total=sum(r.estimated_amount for r in rows))


def test_export_returns_bytes_io():
    result = export_forecast(_make_summary())
    assert isinstance(result, BytesIO)
    assert len(result.read()) > 0


def test_export_has_three_sheets():
    result = export_forecast(_make_summary())
    wb = openpyxl.load_workbook(result)
    assert "預估總覽" in wb.sheetnames
    assert "預估明細" in wb.sheetnames
    assert "排除明細" in wb.sheetnames


def test_export_detail_columns():
    result = export_forecast(_make_summary())
    wb = openpyxl.load_workbook(result)
    ws = wb["預估明細"]
    headers = [cell.value for cell in ws[1]]
    assert "醫院 (客戶)" in headers
    assert "是否排除" in headers


def test_export_excluded_sheet_only_has_excluded_rows():
    rows = [_make_row(excluded=False), _make_row(excluded=True)]
    result = export_forecast(_make_summary(rows=rows))
    wb = openpyxl.load_workbook(result)
    ws = wb["排除明細"]
    data_rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if any(v is not None for v in r)]
    assert len(data_rows) == 1


def test_prepare_export_summary_applies_adjustments_reasons_and_totals():
    from src.backend.forecast_export_workflow import prepare_export_summary

    summary = _make_summary()

    prepared = prepare_export_summary(
        summary,
        manual_adjustments={"A__P1": 15},
        adjustment_reasons={"A__P1": "Review"},
    )

    row = prepared.rows[0]
    assert row.manual_adjustment == 15
    assert row.final_forecast == 15
    assert row.adjustment_reason == "Review"
    assert row.estimated_amount == 1500
    assert prepared.total == 1500
