from io import BytesIO
from pathlib import Path
import re
import shutil
import sqlite3
import uuid

import pandas as pd
from openpyxl import load_workbook

from src.backend import app
from src.backend.forecast_config import ForecastConfig


def _isolated_db_base() -> Path:
    base = Path.cwd() / ".test-dbs" / f"mor-test-{uuid.uuid4().hex}"
    if base.exists():
        shutil.rmtree(base)
    base.mkdir(parents=True)
    return base


def _client(config: dict | None = None):
    app_config = {"TESTING": True, "DB_BASE_PATH": _isolated_db_base()}
    app_config.update(config or {})
    return app.create_app(app_config).test_client()


def _client_with_sales(config: dict | None = None):
    """Create a test client pre-seeded with enough sales data to produce forecast rows."""
    from src.backend.database import get_db
    db_base = _isolated_db_base()
    db = get_db(db_base)
    with db.get_connection() as conn:
        for m, d in [(3, 10), (3, 25), (4, 12), (4, 28)]:
            conn.execute(
                "INSERT INTO sales_records "
                "(order_date, customer_name, product_code, product_name, quantity, unit_price, amount) "
                "VALUES (?, 'A客戶', 'P1', '商品A', 20, 100.0, 2000.0)",
                (f"2026-{m:02d}-{d:02d}",),
            )
        conn.commit()
    app_config = {"TESTING": True, "DB_BASE_PATH": db_base}
    app_config.update(config or {})
    return app.create_app(app_config).test_client()


def _seed_monthly_review_route_data(db_base: Path) -> None:
    from src.backend.database import get_db

    db = get_db(db_base)

    def insert_actuals(year: int, month: int, rows: list[dict]) -> None:
        with db.get_connection() as conn:
            conn.execute(
                """INSERT INTO daily_import_batches
                (source_filename, source_hash, sales_year, sales_month,
                 row_count, quantity_total, taxed_amount_total, status)
                VALUES (?, ?, ?, ?, ?, ?, ?, 'success')""",
                (
                    f"{year}-{month:02d}.xlsx", f"route-{year}-{month}",
                    year, month, len(rows),
                    sum(r["qty"] for r in rows),
                    sum(r["amount"] for r in rows),
                ),
            )
            batch_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
            for row in rows:
                conn.execute(
                    """INSERT INTO daily_sales_actuals
                    (sales_year, sales_month, sales_date, customer_name, product_code,
                     actual_quantity, taxed_amount, bonus_basis_amount, net_unit_price,
                     import_batch_id, customer_code, product_name, sales_quantity, gift_quantity)
                    VALUES (?, ?, ?, ?, ?, ?, ?, 0, ?, ?, '', ?, ?, 0)""",
                    (
                        year, month, f"{year}-{month:02d}-10",
                        row["customer"], row["product"], row["qty"], row["amount"],
                        row["amount"] / row["qty"], batch_id, row["name"], row["qty"],
                    ),
                )
            conn.commit()

    def close_month(year: int, month: int, snapshot_rows: list[dict]) -> None:
        with db.get_connection() as conn:
            cursor = conn.execute(
                """INSERT INTO forecast_snapshots
                (snapshot_name, snapshot_type, year, month, created_by)
                VALUES (?, 'CloseMonth', ?, ?, 'test')""",
                (f"close {year}/{month:02d}", year, month),
            )
            snapshot_id = cursor.lastrowid
            for row in snapshot_rows:
                conn.execute(
                    """INSERT INTO snapshot_items
                    (snapshot_id, customer_name, product_code, system_forecast, final_forecast)
                    VALUES (?, ?, ?, ?, ?)""",
                    (snapshot_id, row["customer"], row["product"], row["forecast"], row["forecast"]),
                )
            conn.execute(
                """INSERT INTO month_close_records
                (year, month, actual_row_count, actual_quantity_total,
                 actual_amount_total, final_snapshot_id)
                VALUES (?, ?, ?, ?, ?, ?)""",
                (
                    year, month, len(snapshot_rows),
                    sum(r["forecast"] for r in snapshot_rows),
                    0.0, snapshot_id,
                ),
            )
            conn.commit()

    insert_actuals(2026, 3, [
        {"customer": "Hospital A", "product": "P1", "name": "Product One", "qty": 300, "amount": 30000},
    ])
    close_month(2026, 3, [{"customer": "Hospital A", "product": "P1", "forecast": 280}])
    insert_actuals(2026, 4, [
        {"customer": "Hospital A", "product": "P1", "name": "Product One", "qty": 200, "amount": 20000},
    ])
    close_month(2026, 4, [{"customer": "Hospital A", "product": "P1", "forecast": 210}])
    insert_actuals(2025, 5, [
        {"customer": "Hospital A", "product": "P1", "name": "Product One", "qty": 120, "amount": 12000},
        {"customer": "Dormant C", "product": "P4", "name": "Product Four", "qty": 40, "amount": 4000},
    ])
    close_month(2025, 5, [
        {"customer": "Hospital A", "product": "P1", "forecast": 120},
        {"customer": "Dormant C", "product": "P4", "forecast": 40},
    ])
    insert_actuals(2026, 5, [
        {"customer": "Hospital A", "product": "P1", "name": "Product One", "qty": 80, "amount": 8000},
        {"customer": "Hospital A", "product": "P2", "name": "Product Two", "qty": 20, "amount": 2000},
        {"customer": "Clinic B", "product": "P3", "name": "Product Three", "qty": 50, "amount": 5000},
    ])
    close_month(2026, 5, [
        {"customer": "Hospital A", "product": "P1", "forecast": 90},
        {"customer": "Hospital A", "product": "P2", "forecast": 10},
        {"customer": "Clinic B", "product": "P3", "forecast": 40},
    ])
    with db.get_connection() as conn:
        conn.execute(
            """INSERT INTO budget_targets
            (year, month, customer_name, product_code, target_quantity, target_amount)
            VALUES (2026, 5, 'Hospital A', 'P1', 100, 12000)"""
        )
        conn.execute(
            """INSERT INTO budget_targets
            (year, month, customer_name, product_code, target_quantity, target_amount)
            VALUES (2026, 5, 'Clinic B', 'P3', 50, 5000)"""
        )
        conn.commit()


def _seed_sales_from_legacy_df(db_base: Path, data: pd.DataFrame) -> None:
    """Seed sales_records from a Chinese-column DataFrame (converts legacy test data to DB rows)."""
    from src.backend.database import get_db
    db = get_db(db_base)
    with db.get_connection() as conn:
        for _, row in data.iterrows():
            year, month, day = int(row["年"]), int(row["月"]), int(row["日"])
            conn.execute(
                "INSERT INTO sales_records "
                "(order_date, customer_name, product_code, product_name, quantity, unit_price, amount) "
                "VALUES (?, ?, ?, ?, ?, ?, ?)",
                (
                    f"{year:04d}-{month:02d}-{day:02d}",
                    str(row["客戶簡稱"]), str(row["商品號"]),
                    str(row.get("商品簡稱", "")),
                    float(row["銷+贈S量"]), float(row["單價NT(淨)"]), float(row["含稅總額(淨)"]),
                ),
            )
        conn.commit()


def test_homepage_loads_dashboard():
    client = _client()

    response = client.get("/")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert "<title>MOR" in html
    assert "<h1>MOR</h1>" in html
    assert '<h1>MOR <span class="app-subtitle">' not in html
    assert "業績總覽" in html
    assert '<a href="/" aria-current="page">業績總覽</a>' in html
    assert "預算目標" in html
    assert "預估達成" in html
    assert "預估業績金額" in html
    assert "金額 GAP" in html
    assert "高風險追蹤" in html
    assert "metrics.target_quantity" not in html
    assert 'href="/forecast"' in html


def test_dashboard_reuses_forecast_context_when_excel_mtime_is_unchanged(monkeypatch):
    calls = {"count": 0}
    real_builder = app.build_forecast_page_context

    def counting_builder(*args, **kwargs):
        calls["count"] += 1
        return real_builder(*args, **kwargs)

    monkeypatch.setattr(app, "build_forecast_page_context", counting_builder)
    client = _client()

    first = client.get("/")
    second = client.get("/")

    assert first.status_code == 200
    assert second.status_code == 200
    assert calls["count"] == 1


def test_forecast_reuses_dashboard_context_cache(monkeypatch):
    calls = {"count": 0}
    real_builder = app.build_forecast_page_context

    def counting_builder(*args, **kwargs):
        calls["count"] += 1
        return real_builder(*args, **kwargs)

    monkeypatch.setattr(app, "build_forecast_page_context", counting_builder)
    client = _client()

    dashboard_response = client.get("/")
    forecast_response = client.get("/forecast")

    assert dashboard_response.status_code == 200
    assert forecast_response.status_code == 200
    assert calls["count"] == 1


def test_save_snapshot_reuses_forecast_context_cache(monkeypatch):
    calls = {"count": 0}
    real_builder = app.build_forecast_page_context

    def counting_builder(*args, **kwargs):
        calls["count"] += 1
        return real_builder(*args, **kwargs)

    db_base_path = _isolated_db_base()
    _seed_db_sales(db_base_path, _minimal_sales_rows())
    monkeypatch.setattr(app, "build_forecast_page_context", counting_builder)
    client = _client({"DB_BASE_PATH": db_base_path})

    forecast_response = client.get("/forecast?year=2026&month=5")
    save_response = client.post(
        "/snapshots/save",
        data={"year": "2026", "month": "5", "snapshot_name": "Cache test"},
    )

    assert forecast_response.status_code == 200
    assert save_response.status_code == 302
    assert calls["count"] == 1


def test_close_month_reuses_forecast_context_cache(monkeypatch):
    calls = {"count": 0}
    real_builder = app.build_forecast_page_context

    def counting_builder(*args, **kwargs):
        calls["count"] += 1
        return real_builder(*args, **kwargs)

    db_base_path = _isolated_db_base()
    _seed_sales_from_legacy_df(
        db_base_path,
        pd.DataFrame([
            {
                "年": 2025, "月": 5, "日": 10,
                "客戶簡稱": "Hospital A", "商品號": "P1",
                "商品簡稱": "Product One", "銷+贈S量": 10,
                "單價NT(淨)": 100, "含稅總額(淨)": 0,
            }
        ]),
    )
    monkeypatch.setattr(app, "build_forecast_page_context", counting_builder)
    client = _client({"DB_BASE_PATH": db_base_path})

    client.post(
        "/monitor/products/import",
        data={"daily_sales_file": (_daily_import_workbook(), "may.xlsx")},
        content_type="multipart/form-data",
    )
    forecast_response = client.get("/forecast?year=2026&month=5")
    close_response = client.post(
        "/monitor/products/close-month",
        data={"year": "2026", "month": "5"},
    )

    assert forecast_response.status_code == 200
    assert close_response.status_code == 302
    assert calls["count"] == 1


def test_homepage_data_health_alert_appears_before_progress_hero():
    db_base_path = _isolated_db_base()
    data = pd.DataFrame([{
        "年": 2026, "月": 4, "日": 10,
        "客戶簡稱": "Hospital A", "商品號": "P1", "商品簡稱": "Product One",
        "銷+贈S量": 3, "單價NT(淨)": 0, "含稅總額(淨)": 0,
    }])
    _seed_sales_from_legacy_df(db_base_path, data)
    client = _client({"DB_BASE_PATH": db_base_path})

    response = client.get("/")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert 'class="alert alert-info' in html
    assert html.index('class="alert alert-info') < html.index('class="hero__container workbench-panel"')


def test_forecast_page_renders_forecast_review_assets_and_tools():
    client = _client_with_sales()

    response = client.get("/forecast")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert 'href="/static/css/mor.css"' in html
    assert 'src="/static/js/forecast-table.js"' in html
    assert 'class="forecast-tools workbench-toolbar"' in html
    assert "總列數" in html
    assert "顯示列數" not in html
    assert "搜尋客戶或品項" in html
    assert "狀態篩選" in html
    assert "data-filter-search" in html
    assert "data-filter-status" in html
    assert "data-edited-count" in html
    assert "data-row-id=" in html
    assert "data-status=" in html
    assert "data-search=" in html


def test_monthly_review_defaults_to_latest_closed_sales_month(monkeypatch):
    from src.backend.monthly_review import MonthlyReviewSummary

    stub_summary = MonthlyReviewSummary(
        year=2026, month=4, closed_at="2026-05-01T00:00:00",
        actual_quantity_total=8.0, forecast_quantity_total=8.0,
        budget_quantity_total=8.0, last_year_quantity_total=6.0,
        actual_amount_total=840.0, forecast_amount_total=840.0,
        budget_amount_total=840.0, last_year_amount_total=567.0,
        forecast_accuracy_total=1.0, yoy_growth_total=1.33,
        budget_achievement_total=1.0,
        forecast_amount_accuracy_total=1.0,
        yoy_amount_growth_total=1.48,
        budget_amount_achievement_total=1.0,
        rows=[],
    )
    from src.backend.monthly_review_context import MonthlyReviewContext

    monkeypatch.setattr(app, "list_reviewable_months", lambda db: [(2026, 4), (2026, 3)])
    monkeypatch.setattr(
        app,
        "build_monthly_review_context",
        lambda db, y, m: MonthlyReviewContext(
            summary=stub_summary,
            action_lists=None,
            customer_summary=None,
            product_summary=None,
            forecast_bias=None,
            trend=None,
            trend_chart=None,
        ),
    )

    client = _client()
    response = client.get("/monthly-review")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert '<a href="/monthly-review" aria-current="page">月底檢討</a>' in html
    assert "2026/04" in html   # latest month shown in picker
    assert "840" in html       # actual_amount_total (含稅淨額) rendered in overview cards


def test_monthly_review_renders_real_sections_with_closed_month_data():
    db_base_path = _isolated_db_base()
    _seed_monthly_review_route_data(db_base_path)
    client = _client({"DB_BASE_PATH": db_base_path})

    response = client.get("/monthly-review?year=2026&month=5")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert "月度總覽（金額為含稅淨額）" in html
    assert "近 12 個月趨勢（金額）" in html
    assert '<svg class="review-trend-svg"' in html
    assert 'points="' in html
    assert "行動清單" in html
    assert "本月優先處理" in html
    assert "review-priority-board" in html
    assert "老闆報告摘要" in html
    assert "預估模型檢討" in html
    assert "客戶面總覽" in html
    assert "品項 Top" in html
    assert "預估準確度" in html
    assert "完整明細" in html
    assert "review-detail-filter" in html
    assert "review-detail-search" in html
    assert "Hospital A" in html
    assert "Product One" in html
    assert "Clinic B" in html
    assert "Dormant C" in html
    assert "準確率" in html
    assert "達成率" in html
    assert "YoY" in html
    assert "66.7%" in html
    assert "83.3%" in html
    assert '本月失準 <span class="badge status-high">2</span>' in html
    assert 'data-search="hospital a p1 product one"' in html
    assert "data-anomaly=\"1\"" in html
    assert 'src="/static/js/monthly-review.js"' in html


def test_frontend_pages_load_htmx_assets():
    client = _client()

    dashboard = client.get("/").get_data(as_text=True)
    forecast = client.get("/forecast").get_data(as_text=True)

    assert "https://unpkg.com/htmx.org@1.9.10" in dashboard
    assert "https://unpkg.com/htmx.org@1.9.10" in forecast


def test_frontend_pages_do_not_load_google_inter_font():
    client = _client()

    pages = [
        client.get("/").get_data(as_text=True),
        client.get("/forecast").get_data(as_text=True),
        client.get("/monitor/products").get_data(as_text=True),
        client.get("/settings").get_data(as_text=True),
    ]
    rendered = "\n".join(pages)

    assert "fonts.googleapis.com" not in rendered
    assert "fonts.gstatic.com" not in rendered
    assert "Inter:wght" not in rendered


def test_dashboard_metrics_partial_renders_fragment_only():
    client = _client()

    response = client.get("/dashboard/metrics")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert "<html" not in html
    assert 'class="hero__container workbench-panel"' in html
    assert "業績總覽" in html


def test_dashboard_copy_does_not_render_mojibake():
    client = _client()

    html = client.get("/").get_data(as_text=True)
    metrics = client.get("/dashboard/metrics").get_data(as_text=True)
    rendered = html + metrics

    assert "業績總覽" in rendered
    assert "預估調整" in rendered
    assert "產品跳單監控" in rendered
    assert "系統設定" in rendered
    assert "年度" in rendered
    assert "月份" in rendered
    assert "套用" in rendered
    assert "預估達成" in rendered
    assert "實績金額" in rendered
    assert "高風險產品" in rendered
    assert "跳單狀態" in rendered
    for broken in ("璆剔蜀蝮質汗", "摰Ｘ", "憸券", "擃", "??/", "?", "?"):
        assert broken not in rendered


def test_rendered_pages_do_not_show_mojibake():
    client = _client_with_sales()

    forecast = client.get("/forecast").get_data(as_text=True)
    product_monitor = client.get("/monitor/products").get_data(as_text=True)
    settings = client.get("/settings").get_data(as_text=True)
    page = client.get("/forecast").get_data(as_text=True)
    row_id = re.search(r'data-row-id="([^"]+)"', page).group(1)
    row_fragment = client.patch(
        f"/forecast/row/{row_id}",
        data={"qty": "", "note": "", "year": "2026", "month": "5"},
    ).get_data(as_text=True)
    rendered = forecast + product_monitor + settings + row_fragment

    assert "自動預估" in rendered
    assert "未到期" in rendered
    assert "還原系統預估" in rendered
    assert "原因..." in rendered
    assert "排除" in rendered
    assert "本月預算為 0" in rendered
    assert "納入" in rendered
    for broken in ("??/", "?祆", "?芸", "?", "?", "頝喳", "蝯梢", ""):
        assert broken not in rendered


def test_dashboard_period_inputs_target_metrics_zone():
    client = _client()

    response = client.get("/")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert 'id="metrics-zone"' in html
    assert 'hx-get="/dashboard/metrics"' in html
    assert 'hx-target="#metrics-zone"' in html


def test_patch_forecast_row_updates_adjustment_and_returns_row_fragment():
    client = _client_with_sales()
    page = client.get("/forecast").get_data(as_text=True)
    row_id = re.search(r'data-row-id="([^"]+)"', page).group(1)

    response = client.patch(
        f"/forecast/row/{row_id}",
        data={"qty": "7.5", "note": "reviewed", "year": "2026", "month": "5"},
    )
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert html.lstrip().startswith("<tr")
    assert f'id="row-{row_id}"' in html
    assert 'value="7"' in html
    assert 'value="reviewed"' in html


def test_forecast_page_exposes_customer_filter_and_row_metadata():
    client = _client_with_sales()

    response = client.get("/forecast")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert 'id="customer-filter"' in html
    # customer filter now uses data-filter-customer (picked up by bindForecastTable in forecast-table.js)
    # instead of the old onchange="filterRows(this.value)" inline handler
    assert 'data-filter-customer' in html
    assert 'data-customer="' in html
    assert 'data-risk="' in html


def test_forecast_page_exposes_view_toggle_for_risk_rows():
    client = _client()

    response = client.get("/forecast")
    html = response.get_data(as_text=True)
    css = Path("static/css/mor.css").read_text(encoding="utf-8")

    assert response.status_code == 200
    assert 'id="btn-anomaly-only"' in html
    assert 'id="btn-show-all"' in html
    # collapsed class is applied client-side by applyViewMode()
    js = Path("static/js/forecast-table.js").read_text(encoding="utf-8")
    assert "forecast-table__row--collapsed" in js
    assert ".forecast-table__row--collapsed" in css


def test_header_navigation_is_consistent_across_frontend_pages():
    db_base_path = _isolated_db_base()
    data = pd.DataFrame([{
        "年": 2026, "月": 4, "日": 10,
        "客戶簡稱": "Hospital A", "商品號": "P1", "商品簡稱": "Product One",
        "銷+贈S量": 3, "單價NT(淨)": 100, "含稅總額(淨)": 0,
    }])
    _seed_sales_from_legacy_df(db_base_path, data)
    client = _client({"DB_BASE_PATH": db_base_path})

    expectations = {
        "/": 'href="/" aria-current="page"',
        "/forecast": 'href="/forecast" aria-current="page"',
        "/monthly-review": 'href="/monthly-review" aria-current="page"',
        "/monitor/products": 'href="/monitor/products" aria-current="page"',
        "/settings": 'href="/settings" aria-current="page"',
    }
    for path, active_link in expectations.items():
        response = client.get(path)
        html = response.get_data(as_text=True)

        assert response.status_code == 200
        assert 'href="/"' in html
        assert 'href="/forecast"' in html
        assert 'href="/monitor/products"' in html
        assert 'href="/settings"' in html
        assert 'href="/exclusions"' not in html
        assert active_link in html
        assert "同步資料" in html


def test_workbench_toolbars_use_shared_structure_classes():
    client = _client()

    pages = {
        "/forecast": client.get("/forecast").get_data(as_text=True),
        "/monitor/products": client.get("/monitor/products").get_data(as_text=True),
        "/settings": client.get("/settings").get_data(as_text=True),
    }
    css = Path("static/css/mor.css").read_text(encoding="utf-8")

    for html in pages.values():
        assert "workbench-toolbar" in html
        assert "toolbar-title" in html
        assert "toolbar-controls" in html

    assert "toolbar-actions" in pages["/forecast"]
    assert "toolbar-actions" in pages["/settings"]
    assert "data-filter-search" in pages["/forecast"]
    assert "data-monitor-search" in pages["/monitor/products"]
    assert "data-item-search" in pages["/settings"]
    assert ".workbench-toolbar" in css
    assert ".toolbar-title" in css
    assert ".toolbar-controls" in css
    assert ".toolbar-actions" in css


def test_workbench_panels_and_tables_use_shared_container_classes():
    client = _client()

    dashboard = client.get("/").get_data(as_text=True)
    forecast = client.get("/forecast").get_data(as_text=True)
    monitor = client.get("/monitor/products").get_data(as_text=True)
    settings = client.get("/settings").get_data(as_text=True)
    css = Path("static/css/mor.css").read_text(encoding="utf-8")

    assert 'class="hero__container workbench-panel' in dashboard
    assert "risk-panel__summary panel workbench-panel" in dashboard
    assert "risk-panel__customer-rank panel workbench-panel" in dashboard
    assert 'class="panel workbench-panel"' in dashboard
    assert "table-wrap compact-table workbench-table-shell" in dashboard

    assert "table-wrap forecast-table-shell workbench-table-shell" in forecast
    assert "monitor-workspace workbench-panel" in monitor
    assert "table-wrap monitor-table-wrap workbench-table-shell" in monitor
    assert "settings-workspace workbench-panel" in settings
    assert "table-wrap settings-table-wrap workbench-table-shell" in settings

    assert ".workbench-panel" in css
    assert ".workbench-table-shell" in css


def test_templates_use_shared_head_assets_and_no_static_inline_layout():
    page_templates = [
        Path("templates/index.html"),
        Path("templates/forecast.html"),
        Path("templates/product_monitor.html"),
        Path("templates/monthly_review.html"),
        Path("templates/settings.html"),
        Path("templates/items.html"),
    ]
    for template_path in page_templates:
        html = template_path.read_text(encoding="utf-8")
        assert '{% include "_head_assets.html" %}' in html
        assert "https://unpkg.com/htmx.org" not in html
        assert "css/mor.css" not in html

    dashboard_partial = Path("templates/_dashboard_metrics.html").read_text(encoding="utf-8")
    head_assets = Path("templates/_head_assets.html").read_text(encoding="utf-8")
    items_template = Path("templates/items.html").read_text(encoding="utf-8")
    css = Path("static/css/mor.css").read_text(encoding="utf-8")

    assert "https://unpkg.com/htmx.org@1.9.10/dist/htmx.min.js" in head_assets
    assert "D1Kt99CQMDuVetoL1lrYwg5t+9QdHe7NLX/SoJYkXDFfX37iInKRy5xLSi8nO7UC" in head_assets
    assert 'style="color:var(--accent-2)"' not in dashboard_partial
    assert 'style="margin:var(--sp-4) 0 0;"' not in dashboard_partial
    assert 'style="display:flex;gap:8px;align-items:center;"' not in dashboard_partial
    assert 'style="color: var(--muted);"' not in items_template

    assert "dist-count--caution" in dashboard_partial
    assert "empty-note" in dashboard_partial
    assert "section-actions" in dashboard_partial
    assert "item-code" in items_template
    assert ".dist-count--caution" in css
    assert ".empty-note" in css
    assert ".section-actions" in css
    assert ".item-code" in css


def test_shared_confirm_dialog_assets_replace_inline_confirm_handlers():
    head_assets = Path("templates/_head_assets.html").read_text(encoding="utf-8")
    header = Path("templates/_header.html").read_text(encoding="utf-8")
    forecast = Path("templates/forecast.html").read_text(encoding="utf-8")
    monitor = Path("templates/product_monitor.html").read_text(encoding="utf-8")
    forecast_js = Path("static/js/forecast-table.js").read_text(encoding="utf-8")
    css = Path("static/css/mor.css").read_text(encoding="utf-8")

    assert "js/ui-feedback.js" in head_assets
    assert '<dialog id="app-confirm" class="app-confirm">' in header
    assert "confirm(" not in header
    assert "confirm(" not in forecast
    assert "confirm(" not in monitor
    assert "window.appConfirm" in forecast_js

    assert 'data-confirm-title="同步 Excel 資料"' in header
    assert 'data-confirm-title="刪除草稿"' in forecast
    assert 'aria-label="刪除草稿"' in forecast
    assert 'id="rd-close" title="關閉" aria-label="關閉詳細資料"' in forecast
    assert 'data-confirm-ok="結月"' in monitor
    assert ".app-confirm" in css
    assert "button.is-loading" in css


def test_frontend_accessibility_tokens_follow_design_roadmap():
    css = Path("static/css/mor.css").read_text(encoding="utf-8")
    design = Path("DESIGN.md").read_text(encoding="utf-8")

    assert 'font-family: "Microsoft JhengHei", "PingFang TC", "Noto Sans TC", "Segoe UI", system-ui, Arial, sans-serif;' in css
    assert "--focus:            #0d9488;" in css
    assert "--focus-ring:       rgba(13, 148, 136, 0.25);" in css
    assert "| `--focus` | `#0d9488` |" in design
    assert "| `--focus-ring` | `rgba(13,148,136,0.25)` |" in design


def test_forecast_empty_state_scroll_and_mid_width_nav_contract():
    forecast = Path("templates/forecast.html").read_text(encoding="utf-8")
    css = Path("static/css/mor.css").read_text(encoding="utf-8")

    assert "請先同步 Excel 資料，或切換到有資料的月份。" in forecast
    assert 'action="/sync"' in forecast
    assert 'data-confirm-title="同步 Excel 資料"' in forecast
    assert "addEventListener('wheel'" not in forecast
    assert "summary.classList.toggle('summary--collapsed', tableWrap.scrollTop > 24)" in forecast
    assert "@media (max-width: 1100px)" in css
    assert ".app-nav::-webkit-scrollbar" in css
    assert "scrollbar-width: none;" in css


def test_old_exclusions_page_redirects_to_item_management():
    db_base_path = _isolated_db_base()
    data = pd.DataFrame([
        {"年": 2026, "月": 4, "日": 10, "客戶簡稱": "Hospital A", "商品號": "P1",
         "商品簡稱": "Product One", "銷+贈S量": 3, "單價NT(淨)": 100, "含稅總額(淨)": 0},
        {"年": 2026, "月": 4, "日": 10, "客戶簡稱": "Hospital A", "商品號": "P2",
         "商品簡稱": "Product Two", "銷+贈S量": 5, "單價NT(淨)": 200, "含稅總額(淨)": 0},
    ])
    _seed_sales_from_legacy_df(db_base_path, data)
    client = _client({"DB_BASE_PATH": db_base_path})

    response = client.get("/exclusions")

    assert response.status_code == 302
    assert response.headers["Location"].endswith("/settings")


def test_item_management_exclusion_is_reflected_on_workbench():
    db_base_path = _isolated_db_base()
    data = pd.DataFrame([
        {"年": 2026, "月": 4, "日": 10, "客戶簡稱": "Hospital A", "商品號": "P1",
         "商品簡稱": "Product One", "銷+贈S量": 3, "單價NT(淨)": 100, "含稅總額(淨)": 0},
        {"年": 2026, "月": 4, "日": 10, "客戶簡稱": "Hospital A", "商品號": "P2",
         "商品簡稱": "Product Two", "銷+贈S量": 5, "單價NT(淨)": 200, "含稅總額(淨)": 0},
    ])
    _seed_sales_from_legacy_df(db_base_path, data)
    client = _client({"DB_BASE_PATH": db_base_path})

    response = client.post(
        "/items/save",
        data={
            "product_codes": ["P1", "P2"],
            "is_budgeted_P1": "1",
            "is_visible_P1": "1",
            "is_budgeted_P2": "1",
            "is_visible_P2": "1",
            "is_excluded_P2": "1",
        },
        follow_redirects=True,
    )
    workbench = client.get("/forecast").get_data(as_text=True)

    assert response.status_code == 200
    assert re.search(r'data-search="[^"]*Product Two"[^>]*data-excluded="true"', workbench)


def test_forecast_page_layers_discontinued_items_below_active_rows():
    db_base_path = _isolated_db_base()
    data = pd.DataFrame([
        {"年": y, "月": m, "日": 10, "客戶簡稱": c, "商品號": p, "商品簡稱": n,
         "銷+贈S量": q, "單價NT(淨)": pr, "含稅總額(淨)": 0}
        for y, m, c, p, n, q, pr in (
            (2026, 2, "Hospital A", "P1", "Product Active", 10, 100),
            (2026, 3, "Hospital A", "P1", "Product Active", 10, 100),
            (2026, 4, "Hospital A", "P1", "Product Active", 10, 100),
            (2025, 5, "Hospital A", "P2", "Product Old", 100, 100),
        )
    ])
    _seed_sales_from_legacy_df(db_base_path, data)
    client = _client({"DB_BASE_PATH": db_base_path})

    response = client.post(
        "/items/save",
        data={
            "product_codes": ["P1", "P2"],
            "is_budgeted_P1": "1",
            "is_visible_P1": "1",
            "item_status_P1": "active",
            "is_budgeted_P2": "1",
            "is_visible_P2": "1",
            "item_status_P2": "discontinued",
        },
        follow_redirects=True,
    )
    forecast = client.get("/forecast?year=2026&month=5").get_data(as_text=True)

    assert response.status_code == 200
    assert "已停用品項" in forecast
    assert "共 1 項" in forecast
    assert "去年業績合計" in forecast
    assert "10,000" in forecast
    assert 'data-item-status="active"' in forecast
    assert 'data-item-status="discontinued"' in forecast
    assert forecast.index("Product Active") < forecast.index("已停用品項") < forecast.index("Product Old")


def test_settings_page_uses_item_status_without_auxiliary_labels():
    db_base_path = _isolated_db_base()
    data = pd.DataFrame([{
        "年": 2026, "月": 4, "日": 10,
        "客戶簡稱": "Hospital A", "商品號": "P1", "商品簡稱": "Product One",
        "銷+贈S量": 3, "單價NT(淨)": 100, "含稅總額(淨)": 0,
    }])
    _seed_sales_from_legacy_df(db_base_path, data)
    client = _client({"DB_BASE_PATH": db_base_path})

    settings = client.get("/settings").get_data(as_text=True)

    assert 'name="item_status_P1"' in settings
    assert "使用中" in settings
    assert "停用" in settings
    assert "特殊品項" not in settings
    assert "新上市" not in settings
    assert "status_label_P1" not in settings


def test_item_settings_save_price_quantity_and_forecast_uses_it():
    db_base_path = _isolated_db_base()
    data = pd.DataFrame([
        {"年": 2026, "月": m, "日": 10, "客戶簡稱": "Hospital A", "商品號": "P1",
         "商品簡稱": "Product One", "銷+贈S量": 300, "單價NT(淨)": 3200, "含稅總額(淨)": 0}
        for m in (1, 2, 3)
    ])
    _seed_sales_from_legacy_df(db_base_path, data)
    client = _client({"DB_BASE_PATH": db_base_path})
    with sqlite3.connect(db_base_path / "mor_workbench.db") as conn:
        conn.execute(
            """
            INSERT INTO budget_targets
            (year, month, customer_name, product_code, target_quantity, target_amount, base_target_quantity)
            VALUES (2026, 4, 'Hospital A', 'P1', 300, 3000, 1)
            """
        )
        conn.commit()

    settings = client.get("/settings?year=2026&month=4").get_data(as_text=True)
    assert 'name="price_quantity_P1"' in settings

    response = client.post(
        "/items/save",
        data={
            "product_codes": ["P1"],
            "is_budgeted_P1": "1",
            "is_visible_P1": "1",
            "price_quantity_P1": "100",
        },
        follow_redirects=True,
    )
    forecast = client.get("/forecast?year=2026&month=4").get_data(as_text=True)

    assert response.status_code == 200
    saved_settings = response.get_data(as_text=True)
    assert 'name="price_quantity_P1" type="number" min="0" step="1" value="100"' in saved_settings
    assert 'value="100.0"' not in saved_settings
    with sqlite3.connect(db_base_path / "mor_workbench.db") as conn:
        columns = [row[1] for row in conn.execute("PRAGMA table_info(item_configs)")]
        saved = conn.execute("SELECT price_quantity FROM item_configs WHERE product_code = 'P1'").fetchone()
    assert "price_quantity" in columns
    assert saved == (100,)
    assert 'data-price-quantity="100' in forecast
    # Estimated amount is now computed client-side from data attributes
    assert 'data-price=' in forecast


def test_item_settings_save_persists_payload_shape_to_item_configs():
    db_base_path = _isolated_db_base()
    from src.backend.database import get_db
    db = get_db(db_base_path)
    with db.get_connection() as conn:
        conn.execute(
            """
            INSERT INTO sales_records
            (order_date, customer_name, product_code, product_name, quantity, unit_price, amount)
            VALUES ('2026-04-10', 'Hospital A', 'P1', 'Product One', 3, 100, 300)
            """
        )
        conn.commit()
    client = _client({"DB_BASE_PATH": db_base_path})

    response = client.post(
        "/items/save",
        data={
            "product_codes": ["P1"],
            "is_excluded_P1": "1",
            "is_visible_P1": "1",
            "price_quantity_P1": "24.7",
            "item_status_P1": "discontinued",
        },
        follow_redirects=True,
    )

    assert response.status_code == 200
    assert "settings-workspace" in response.get_data(as_text=True)
    with sqlite3.connect(db_base_path / "mor_workbench.db") as conn:
        saved = conn.execute(
            """
            SELECT is_excluded, is_budgeted, is_visible, price_quantity, item_status
            FROM item_configs
            WHERE product_code = 'P1'
            """
        ).fetchone()
    assert saved == (1, 0, 1, 24.0, "discontinued")


def test_item_settings_save_coerces_invalid_fields_to_existing_defaults():
    db_base_path = _isolated_db_base()
    from src.backend.database import get_db
    db = get_db(db_base_path)
    with db.get_connection() as conn:
        conn.execute(
            """
            INSERT INTO sales_records
            (order_date, customer_name, product_code, product_name, quantity, unit_price, amount)
            VALUES ('2026-04-10', 'Hospital A', 'P1', 'Product One', 3, 100, 300)
            """
        )
        conn.commit()
    client = _client({"DB_BASE_PATH": db_base_path})

    response = client.post(
        "/items/save",
        data={
            "product_codes": ["P1"],
            "price_quantity_P1": "not-a-number",
            "item_status_P1": "retired",
        },
        follow_redirects=True,
    )

    assert response.status_code == 200
    with sqlite3.connect(db_base_path / "mor_workbench.db") as conn:
        saved = conn.execute(
            """
            SELECT is_excluded, is_budgeted, is_visible, price_quantity, item_status
            FROM item_configs
            WHERE product_code = 'P1'
            """
        ).fetchone()
    assert saved == (0, 0, 0, 0.0, "active")


def test_item_settings_save_clears_forecast_context_cache():
    db_base_path = _isolated_db_base()
    data = pd.DataFrame([
        {"年": 2026, "月": m, "日": 10, "客戶簡稱": "Hospital A", "商品號": "P1",
         "商品簡稱": "Product One", "銷+贈S量": 100, "單價NT(淨)": 1000, "含稅總額(淨)": 0}
        for m in (1, 2, 3)
    ])
    _seed_sales_from_legacy_df(db_base_path, data)
    client = _client({"DB_BASE_PATH": db_base_path})

    first_forecast = client.get("/forecast?year=2026&month=4").get_data(as_text=True)
    response = client.post(
        "/items/save",
        data={
            "product_codes": ["P1"],
            "is_budgeted_P1": "1",
        },
        follow_redirects=True,
    )
    second_forecast = client.get("/forecast?year=2026&month=4").get_data(as_text=True)

    assert "Product One" in first_forecast
    assert response.status_code == 200
    assert "Product One" not in second_forecast


def test_unbudgeted_item_has_no_budget_target_on_forecast_page():
    db_base_path = _isolated_db_base()
    data = pd.DataFrame([
        {"年": 2026, "月": m, "日": 10, "客戶簡稱": "Hospital A", "商品號": "P1",
         "商品簡稱": "Product One", "銷+贈S量": 100, "單價NT(淨)": 1000, "含稅總額(淨)": 0}
        for m in (1, 2, 3)
    ])
    _seed_sales_from_legacy_df(db_base_path, data)
    client = _client({"DB_BASE_PATH": db_base_path})
    with sqlite3.connect(db_base_path / "mor_workbench.db") as conn:
        conn.execute(
            """
            INSERT INTO budget_targets
            (year, month, customer_name, product_code, target_quantity, target_amount, base_target_quantity)
            VALUES (2026, 4, 'Hospital A', 'P1', 300, 3000, 1)
            """
        )
        conn.commit()

    response = client.post(
        "/items/save",
        data={
            "product_codes": ["P1"],
            "is_visible_P1": "1",
        },
        follow_redirects=True,
    )
    forecast = client.get("/forecast?year=2026&month=4").get_data(as_text=True)

    assert response.status_code == 200
    assert 'data-budget="0.0"' in forecast


def test_adjustment_save_migrates_old_database_without_updated_by():
    db_base_path = _isolated_db_base()
    db_path = db_base_path / "mor_workbench.db"
    with sqlite3.connect(db_path) as conn:
        conn.execute(
            """
            CREATE TABLE forecast_adjustments (
                year INTEGER,
                month INTEGER,
                customer_name TEXT,
                product_code TEXT,
                manual_quantity REAL,
                adjustment_reason TEXT,
                updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                PRIMARY KEY (year, month, customer_name, product_code)
            )
            """
        )
        conn.commit()

    client = _client({"DB_BASE_PATH": db_base_path})

    response = client.post(
        "/adjustments/save",
        data={
            "row_id": "Hospital A__P1",
            "manual_adjustment": "12",
            "reason": "Review",
            "year": "2026",
            "month": "5",
        },
    )

    assert response.status_code == 200
    with sqlite3.connect(db_path) as conn:
        columns = [row[1] for row in conn.execute("PRAGMA table_info(forecast_adjustments)")]
        saved = conn.execute(
            """
            SELECT manual_quantity, adjustment_reason, updated_by
            FROM forecast_adjustments
            WHERE year = 2026 AND month = 5 AND customer_name = 'Hospital A' AND product_code = 'P1'
            """
        ).fetchone()
    assert "updated_by" in columns
    assert saved == (12, "Review", "User")


def test_forecast_write_workflow_save_row_override_persists_adjustment():
    from src.backend.database import get_db
    from src.backend.forecast_write_workflow import save_row_override

    db_base_path = _isolated_db_base()
    db = get_db(db_base_path)

    save_row_override(db, "Hospital A__P1", "12.7", "Review", 2026, 5)

    with sqlite3.connect(db_base_path / "mor_workbench.db") as conn:
        saved = conn.execute(
            """
            SELECT manual_quantity, adjustment_reason, updated_by
            FROM forecast_adjustments
            WHERE year = 2026 AND month = 5 AND customer_name = 'Hospital A' AND product_code = 'P1'
            """
        ).fetchone()
    assert saved == (12, "Review", "User")


def test_forecast_page_renders_review_validation_and_accessibility_hooks():
    client = _client_with_sales()

    response = client.get("/forecast")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert 'aria-label="人工調整數量"' in html
    assert re.search(r'name="manual_adjustment__[^"]+" type="number" min="0" step="1"', html)
    assert 'step="0.01"' not in html
    assert 'data-validation-message' in html
    assert "人工數量必須是 0 以上的數字" in html


def test_adjustment_save_manual_quantity_as_integer():
    db_base_path = _isolated_db_base()
    client = _client({"DB_BASE_PATH": db_base_path})

    response = client.post(
        "/adjustments/save",
        data={
            "row_id": "Hospital A__P1",
            "manual_adjustment": "12.7",
            "reason": "Review",
            "year": "2026",
            "month": "5",
        },
    )

    assert response.status_code == 200
    with sqlite3.connect(db_base_path / "mor_workbench.db") as conn:
        saved = conn.execute(
            """
            SELECT manual_quantity
            FROM forecast_adjustments
            WHERE year = 2026 AND month = 5 AND customer_name = 'Hospital A' AND product_code = 'P1'
            """
        ).fetchone()
    assert saved == (12,)


def test_forecast_rows_expose_dashboard_jump_anchor():
    client = _client_with_sales()

    response = client.get("/forecast")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    row_id = re.search(r'data-row-id="([^"]+)"', html).group(1)
    assert f'id="row-{row_id}"' in html


def test_static_assets_define_invalid_row_validation_behavior():
    css = Path("static/css/mor.css").read_text(encoding="utf-8")
    js = Path("static/js/forecast-table.js").read_text(encoding="utf-8")

    assert "tr.invalid" in css
    assert "validateBeforeSubmit" in js
    assert "aria-invalid" in js
    assert "data-unrendered-total" in js


def test_forecast_row_keeps_three_char_abbreviation_except_eli_dose():
    row_template = Path("templates/_forecast_row.html").read_text(encoding="utf-8")
    css = Path("static/css/mor.css").read_text(encoding="utf-8")

    assert "product_display_name" in row_template
    assert "min-width: 96px;" in css


def test_forecast_table_compares_live_gap_to_budget():
    js = Path("static/js/forecast-table.js").read_text(encoding="utf-8")

    assert "const gap = finalQty - budget;" in js
    assert "finalQty - state.actualQuantity" not in js


def test_forecast_table_totals_only_include_budgeted_rows():
    js = Path("static/js/forecast-table.js").read_text(encoding="utf-8")
    row_template = Path("templates/_forecast_row.html").read_text(encoding="utf-8")

    assert "budgetQuantity: Number(row.dataset.budget || 0)" in js
    assert "priceQuantity: Number(row.dataset.priceQuantity || 1)" in js
    assert "(finalForecastQuantity(state) / priceQuantity) * state.price" in js
    assert "state.budgetQuantity <= 0" in js
    assert "data-price-quantity=" in row_template


def test_forecast_detail_uses_neutral_zero_budget_status():
    js = Path("static/js/forecast-table.js").read_text(encoding="utf-8")

    assert "缺預算" in js
    assert "預算為 0" not in js


def test_css_keeps_letter_spacing_neutral_for_dense_operational_ui():
    css = Path("static/css/mor.css").read_text(encoding="utf-8")

    non_zero_letter_spacing = [
        value.strip()
        for value in re.findall(r"letter-spacing:\s*([^;]+);", css)
        if value.strip() not in {"0", "0em", "0px"}
    ]

    assert non_zero_letter_spacing == []


def test_dashboard_css_uses_bem_class_names():
    css = Path("static/css/mor.css").read_text(encoding="utf-8")
    dashboard = Path("templates/_dashboard_metrics.html").read_text(encoding="utf-8")
    for old_selector in (".progress-hero", ".sub-metrics", ".risk-summary", ".customer-rank", ".row-collapsed"):
        assert old_selector not in css
    for old_class in (
        'class="progress-hero',
        'class="sub-metrics',
        'class="risk-summary',
        'class="customer-rank',
        'class="row-collapsed',
    ):
        assert old_class not in dashboard

    assert "hero__container" in css + dashboard
    assert "hero__sub-metrics" in css + dashboard
    assert "risk-panel__summary" in css + dashboard
    assert "risk-panel__customer-rank" in css + dashboard
    assert "forecast-table__row--collapsed" in css


def test_dashboard_sparkline_canvas_ids_do_not_depend_on_entity_labels():
    dashboard = Path("templates/_dashboard_metrics.html").read_text(encoding="utf-8")
    table_module = Path("static/js/analytics-table.js").read_text(encoding="utf-8")

    assert "entity_id}`.replace" not in dashboard
    # Canvas ids are generated by AnalyticsTable from idPrefix + row index, never from labels
    assert 'idPrefix: "customer"' in dashboard
    assert 'idPrefix: "product"' in dashboard
    assert "spark-${opts.idPrefix}-${idx}" in table_module


def test_dashboard_progress_hero_uses_dense_metric_layout():
    css = Path("static/css/mor.css").read_text(encoding="utf-8")

    assert "grid-template-columns: minmax(0, 1fr) minmax(440px, 520px);" in css
    assert ".hero__sub-metrics {\n  display: grid;\n  grid-template-columns: repeat(2, minmax(0, 1fr));" in css


def test_homepage_loads_gracefully_with_empty_db():
    """DB-first: empty DB produces 200 with no hard error (Excel path irrelevant)."""
    client = _client()   # isolated empty DB, DATA_BASE_PATH not used for rendering

    response = client.get("/")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert "業績總覽" in html   # page structure renders even with no data


def test_homepage_exposes_unrendered_total_when_rows_are_limited():
    config = ForecastConfig(visible_row_limit=1)
    db_base_path = _isolated_db_base()
    data = pd.DataFrame([
        {"年": 2026, "月": m, "日": 15, "客戶簡稱": c, "商品號": p,
         "商品簡稱": f"Product {p}", "銷+贈S量": 10, "單價NT(淨)": 100, "含稅總額(淨)": 0}
        for c, p in (("A", "P1"), ("B", "P2"))
        for m in (1, 2, 3)
    ])
    _seed_sales_from_legacy_df(db_base_path, data)
    client = _client({"FORECAST_CONFIG": config, "DB_BASE_PATH": db_base_path})
    with sqlite3.connect(db_base_path / "mor_workbench.db") as conn:
        conn.execute(
            """
            INSERT INTO budget_targets
            (year, month, customer_name, product_code, target_quantity, target_amount, base_target_quantity)
            VALUES (2026, 4, 'B', 'P2', 10, 1000, 10)
            """
        )
        conn.commit()

    response = client.get("/forecast")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert html.count("data-row-id=") == 1
    assert 'data-unrendered-total value="1000.0"' in html
    assert "1 / 2" in html


def test_product_monitor_page_renders_drop_table():
    db_base_path = _isolated_db_base()
    data = pd.DataFrame([
        {"年": 2025, "月": 5, "日": 10, "客戶簡稱": "Hospital A", "商品號": "P1",
         "商品簡稱": "ELI 22.5癌立佳", "銷+贈S量": 100, "單價NT(淨)": 100, "含稅總額(淨)": 0},
        {"年": 2026, "月": 4, "日": 10, "客戶簡稱": "Hospital A", "商品號": "P1",
         "商品簡稱": "ELI 22.5癌立佳", "銷+贈S量": 40, "單價NT(淨)": 100, "含稅總額(淨)": 0},
    ])
    _seed_sales_from_legacy_df(db_base_path, data)
    client = _client({"DB_BASE_PATH": db_base_path})

    response = client.get("/monitor/products?year=2026&month=5")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert "產品跳單監控" in html
    assert "狀態" in html
    assert "去年成長率" in html
    assert "預算達成率" in html
    assert "週期狀態" in html
    assert "高風險" in html
    assert html.index("monitor-sticky--status") < html.index("monitor-sticky--customer")
    assert html.index("monitor-sticky--customer") < html.index("monitor-sticky--product")
    assert 'data-monitor-sort="status"' in html
    assert 'data-monitor-sort="forecast"' in html
    assert 'aria-sort="none"' in html
    assert 'title="ELI 22.5癌立佳">ELI 22.5</td>' in html
    assert "ID: P1" not in html
    assert "data-monitor-search" in html


def test_settings_page_renders_item_config_and_data_checks():
    db_base_path = _isolated_db_base()
    data = pd.DataFrame([{
        "年": 2026, "月": 4, "日": 10,
        "客戶簡稱": "Hospital A", "商品號": "P1", "商品簡稱": "Product One",
        "銷+贈S量": 3, "單價NT(淨)": 0, "含稅總額(淨)": 0,
    }])
    _seed_sales_from_legacy_df(db_base_path, data)
    client = _client({"DB_BASE_PATH": db_base_path})

    response = client.get("/settings")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert "系統設定 / 資料檢核" in html
    assert "排除預估" in html
    assert "預算資料檢查" in html
    assert "訂單資料檢查" in html
    assert "品項設定" in html
    assert "settings-workspace" in html
    assert "settings-health-panel" in html
    assert "儲存設定" in html


def test_dashboard_omits_zero_budget_notice_but_keeps_price_warning():
    db_base_path = _isolated_db_base()
    data = pd.DataFrame([{
        "年": 2026, "月": 4, "日": 10,
        "客戶簡稱": "Hospital A", "商品號": "P1", "商品簡稱": "Product One",
        "銷+贈S量": 3, "單價NT(淨)": 0, "含稅總額(淨)": 0,
    }])
    _seed_sales_from_legacy_df(db_base_path, data)
    client = _client({"DB_BASE_PATH": db_base_path})

    dashboard = client.get("/").get_data(as_text=True)
    settings = client.get("/settings").get_data(as_text=True)

    assert "有 1 筆缺少預算目標" not in dashboard
    assert "有 1 筆單價為 0" in dashboard
    assert "有 1 筆缺少預算目標" not in settings
    assert "有 1 筆單價為 0" in settings  # settings now displays data_issues


def test_items_route_redirects_to_settings():
    client = _client()

    response = client.get("/items")

    assert response.status_code == 302
    assert response.headers["Location"].endswith("/settings")


def test_settings_page_exposes_item_search_tools():
    db_base_path = _isolated_db_base()
    data = pd.DataFrame([{
        "年": 2026, "月": 4, "日": 10,
        "客戶簡稱": "Hospital A", "商品號": "P1", "商品簡稱": "Product One",
        "銷+贈S量": 3, "單價NT(淨)": 100, "含稅總額(淨)": 0,
    }])
    _seed_sales_from_legacy_df(db_base_path, data)
    client = _client({"DB_BASE_PATH": db_base_path})

    response = client.get("/settings")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert 'src="/static/js/item-settings.js"' in html
    assert "data-item-search" in html
    assert "data-item-visible" in html
    assert re.search(r'data-item-row[^>]+data-search="[^"]*Product One[^"]*P1', html)


def test_export_rejects_unknown_review_row_id_without_workbook():
    client = _client()

    response = client.post(
        "/export",
        data={"year": "2026", "month": "5", "manual_adjustment__missing__row": "10"},
    )

    assert response.status_code == 400
    assert "Unknown forecast row" in response.get_data(as_text=True)


def test_export_accepts_current_forecast_signature():
    client = _client()
    review_response = client.get("/forecast?year=2026&month=5")
    signature = re.search(
        r'name="forecast_signature" value="([^"]+)"',
        review_response.get_data(as_text=True),
    ).group(1)

    export_response = client.post(
        "/export",
        data={"year": "2026", "month": "5", "forecast_signature": signature},
    )

    assert export_response.status_code == 200
    assert export_response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_export_workbook_readback_has_expected_tabs_and_totals():
    client = _client_with_sales()
    review_response = client.get("/forecast?year=2026&month=5")
    signature = re.search(
        r'name="forecast_signature" value="([^"]+)"',
        review_response.get_data(as_text=True),
    ).group(1)

    export_response = client.post(
        "/export",
        data={"year": "2026", "month": "5", "forecast_signature": signature},
    )

    assert export_response.status_code == 200
    workbook = load_workbook(BytesIO(export_response.data), data_only=True)
    assert workbook.sheetnames == ["預估總覽", "預估明細", "排除明細"]

    summary_sheet = workbook["預估總覽"]
    summary_values = {
        summary_sheet.cell(row=row_idx, column=1).value: summary_sheet.cell(row=row_idx, column=2).value
        for row_idx in range(2, summary_sheet.max_row + 1)
    }
    detail_sheet = workbook["預估明細"]
    headers = [cell.value for cell in detail_sheet[1]]
    amount_index = headers.index("預估金額")
    excluded_index = headers.index("是否排除")
    detail_rows = list(detail_sheet.iter_rows(min_row=2, values_only=True))
    included_rows = [row for row in detail_rows if row[excluded_index] != "是"]
    excluded_rows = [row for row in detail_rows if row[excluded_index] == "是"]

    assert summary_values["預估月份"] == "2026/05"
    assert summary_values["預估總金額"] == sum((row[amount_index] or 0) for row in included_rows)
    assert summary_values["列入預估品項數"] == len(included_rows)
    assert summary_values["排除品項數"] == len(excluded_rows)


def test_export_rejects_stale_forecast_signature():
    """Stale signature: add a forecast adjustment between review and export to change the forecast."""
    config = ForecastConfig(visible_row_limit=1)
    db_base_path = _isolated_db_base()
    # Seed two customers so there are at least 2 forecast rows
    data = pd.DataFrame([
        {"年": 2026, "月": m, "日": 15, "客戶簡稱": c, "商品號": p,
         "商品簡稱": f"Product {p}", "銷+贈S量": 10, "單價NT(淨)": 100, "含稅總額(淨)": 0}
        for c, p in (("A", "P1"), ("B", "P2"))
        for m in (1, 2, 3)
    ])
    _seed_sales_from_legacy_df(db_base_path, data)
    client = _client({"FORECAST_CONFIG": config, "DB_BASE_PATH": db_base_path})

    # GET /forecast — capture signature
    review_response = client.get("/forecast?year=2026&month=4")
    signature = re.search(
        r'name="forecast_signature" value="([^"]+)"',
        review_response.get_data(as_text=True),
    ).group(1)

    # Mutate B__P2 forecast via adjustment → cache clears → new signature on next build
    client.patch(
        "/forecast/row/B__P2",
        data={"year": "2026", "month": "4", "qty": "99"},
        content_type="application/x-www-form-urlencoded",
    )

    export_response = client.post(
        "/export",
        data={"year": "2026", "month": "4", "forecast_signature": signature},
    )

    assert export_response.status_code == 400
    assert "Forecast review changed" in export_response.get_data(as_text=True)


def test_export_rejects_invalid_manual_quantity_without_500():
    client = _client()

    response = client.post(
        "/export",
        data={"year": "2026", "month": "5", "manual_adjustment__A__P1": "abc"},
    )

    assert response.status_code == 400
    assert "人工數量" in response.get_data(as_text=True)


def test_forecast_table_rebinds_row_events_after_htmx_swap():
    script = Path("static/js/forecast-table.js").read_text(encoding="utf-8")

    assert "htmx:afterSwap" in script
    assert "bindForecastRow" in script
    assert "data-forecast-bound" in script


def test_forecast_table_supports_keyboard_navigation_for_manual_inputs():
    script = Path("static/js/forecast-table.js").read_text(encoding="utf-8")

    assert "function visibleDataRows()" in script
    assert "function bindKeyboardNavigation()" in script
    assert '"Enter"' in script
    assert '"ArrowDown"' in script
    assert '"ArrowUp"' in script
    assert '"Escape"' in script
    assert "e.preventDefault()" in script
    assert "target.focus()" in script
    assert "target.select()" in script
    assert "restore.click()" in script
    assert "bindKeyboardNavigation();" in script


# ---------------------------------------------------------------------------
# Daily sales import route and upload UI
# ---------------------------------------------------------------------------

def _daily_import_workbook() -> BytesIO:
    stream = BytesIO()
    pd.DataFrame(
        [
            {
                "出貨日期": "2026-05-04",
                "客戶代號": "C001",
                "客戶簡稱": "Hospital A",
                "產品": "P1",
                "產品簡稱": "Product One",
                "銷售數量": 5,
                "贈品數量": 1,
                "銷貨淨價": 100,
                "含稅淨額": 600,
                "折後業績": 0,
                "發票編號": "INV",
                "出貨單號": "SHIP",
                "單別": "正常銷",
                "業績屬性": "處方",
            }
        ]
    ).to_excel(stream, index=False)
    stream.seek(0)
    return stream


def test_product_monitor_import_route_stores_daily_actuals():
    db_base_path = _isolated_db_base()
    data = pd.DataFrame([
        {"年": 2025, "月": 5, "日": 10, "客戶簡稱": "Hospital A", "商品號": "P1",
         "商品簡稱": "Product One", "銷+贈S量": 10, "單價NT(淨)": 100, "含稅總額(淨)": 0},
        {"年": 2026, "月": 4, "日": 10, "客戶簡稱": "Hospital A", "商品號": "P1",
         "商品簡稱": "Product One", "銷+贈S量": 3, "單價NT(淨)": 100, "含稅總額(淨)": 0},
    ])
    _seed_sales_from_legacy_df(db_base_path, data)
    client = _client({"DB_BASE_PATH": db_base_path})

    response = client.post(
        "/monitor/products/import",
        data={"daily_sales_file": (_daily_import_workbook(), "daily.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert "匯入完成" in html
    assert "6" in html


def test_product_monitor_page_exposes_daily_sales_import_form():
    client = _client()

    html = client.get("/monitor/products").get_data(as_text=True)

    assert 'action="/monitor/products/import"' in html
    assert 'name="daily_sales_file"' in html
    assert "選擇檔案" in html


def test_close_month_route_creates_record_and_blocks_reimport():
    db_base_path = _isolated_db_base()
    data = pd.DataFrame([{
        "年": 2025, "月": 5, "日": 10, "客戶簡稱": "Hospital A", "商品號": "P1",
        "商品簡稱": "Product One", "銷+贈S量": 10, "單價NT(淨)": 100, "含稅總額(淨)": 0,
    }])
    _seed_sales_from_legacy_df(db_base_path, data)
    client = _client({"DB_BASE_PATH": db_base_path})

    # Import May 2026 data
    client.post(
        "/monitor/products/import",
        data={"daily_sales_file": (_daily_import_workbook(), "may.xlsx")},
        content_type="multipart/form-data",
    )
    # Close May 2026
    close_resp = client.post(
        "/monitor/products/close-month",
        data={"year": "2026", "month": "5"},
        follow_redirects=True,
    )
    assert "結月完成" in close_resp.get_data(as_text=True)

    # Reimport should be blocked
    blocked = client.post(
        "/monitor/products/import",
        data={"daily_sales_file": (_daily_import_workbook(), "may_v2.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert "已結月" in blocked.get_data(as_text=True)


def test_monthly_review_page_loads_without_data():
    client = _client()
    response = client.get("/monthly-review")
    html = response.get_data(as_text=True)
    assert response.status_code == 200
    assert "月底檢討" in html
    assert "尚無結月資料" in html
    assert 'js/monthly-review.js' in html
    assert "querySelectorAll('tbody tr')" not in html


def test_monthly_review_template_uses_section_partials():
    template = Path("templates/monthly_review.html").read_text(encoding="utf-8")
    summary_partial = Path("templates/_monthly_review_summary.html").read_text(encoding="utf-8")
    tables_partial = Path("templates/_monthly_review_tables.html").read_text(encoding="utf-8")
    detail_partial = Path("templates/_monthly_review_detail.html").read_text(encoding="utf-8")
    macros = Path("templates/_value_macros.html").read_text(encoding="utf-8")
    assert '{% include "_monthly_review_summary.html" %}' in template
    assert '{% include "_monthly_review_insights.html" %}' in template
    assert '{% include "_monthly_review_tables.html" %}' in template
    assert '{% include "_monthly_review_detail.html" %}' in template
    assert template.count("workbench-panel review-section") <= 1
    assert "namespace(" not in summary_partial
    assert "chart_w" not in summary_partial
    assert "{% macro yoy(" in macros
    assert "_value_macros.html" in summary_partial
    assert "_value_macros.html" in tables_partial
    assert "_value_macros.html" in detail_partial
    assert "{% set y =" not in tables_partial
    assert "{% set yoy =" not in detail_partial


def test_close_month_auto_saves_forecast_snapshot():
    """結月路由應自動儲存最終預估快照，供 Phase 7 月底檢討頁使用。"""
    db_base_path = _isolated_db_base()
    data = pd.DataFrame([{
        "年": 2025, "月": 5, "日": 10, "客戶簡稱": "Hospital A", "商品號": "P1",
        "商品簡稱": "Product One", "銷+贈S量": 10, "單價NT(淨)": 100, "含稅總額(淨)": 0,
    }])
    _seed_sales_from_legacy_df(db_base_path, data)
    client = _client({"DB_BASE_PATH": db_base_path})

    client.post(
        "/monitor/products/import",
        data={"daily_sales_file": (_daily_import_workbook(), "may.xlsx")},
        content_type="multipart/form-data",
    )
    client.post("/monitor/products/close-month", data={"year": "2026", "month": "5"})

    from src.backend.daily_sales_importer import get_close_record
    from src.backend.database import get_db
    db = get_db(db_base_path)
    rec = get_close_record(db, 2026, 5)

    assert rec is not None
    assert rec["final_snapshot_id"] is not None

    from src.backend.snapshot_service import load_snapshot_items
    items = load_snapshot_items(db, rec["final_snapshot_id"])
    assert len(items) >= 1  # at least one forecast row saved


def test_forecast_page_labels_close_month_snapshot_and_hides_delete_action():
    db_base_path = _isolated_db_base()
    _seed_db_sales(db_base_path, _minimal_sales_rows())

    from src.backend.database import get_db
    from src.backend.snapshot_service import save_snapshot

    db = get_db(db_base_path)
    save_snapshot(db, 2026, 5, "Draft snapshot", "Draft", [])
    save_snapshot(db, 2026, 5, "CloseMonth snapshot", "CloseMonth", [], created_by="close-month")
    client = _client({"DB_BASE_PATH": db_base_path})

    html = client.get("/forecast?year=2026&month=5").get_data(as_text=True)

    assert "CloseMonth" in html
    assert "Draft" in html
    close_month_block = html.split("Draft snapshot", 1)[0]
    draft_block = html.split("Draft snapshot", 1)[1]
    assert "CloseMonth snapshot" in close_month_block
    assert 'action="/snapshots/delete"' not in close_month_block
    assert 'action="/snapshots/delete"' in draft_block


def test_product_monitor_shows_close_button_after_import():
    db_base_path = _isolated_db_base()
    data = pd.DataFrame([{
        "年": 2025, "月": 5, "日": 10, "客戶簡稱": "Hospital A", "商品號": "P1",
        "商品簡稱": "Product One", "銷+贈S量": 10, "單價NT(淨)": 100, "含稅總額(淨)": 0,
    }])
    _seed_sales_from_legacy_df(db_base_path, data)
    client = _client({"DB_BASE_PATH": db_base_path})

    client.post(
        "/monitor/products/import",
        data={"daily_sales_file": (_daily_import_workbook(), "may.xlsx")},
        content_type="multipart/form-data",
    )
    html = client.get("/monitor/products?year=2026&month=5").get_data(as_text=True)

    assert 'action="/monitor/products/close-month"' in html
    assert "結月" in html


# ---------------------------------------------------------------------------
# Route guard: normal page/export routes must NOT read Excel (Phase 3)
# ---------------------------------------------------------------------------

def _seed_db_sales(db_base: Path, rows: list[dict]) -> None:
    """Seed sales_records with minimal rows for guard tests."""
    from src.backend.database import get_db
    db = get_db(db_base)
    with db.get_connection() as conn:
        for row in rows:
            conn.execute(
                "INSERT INTO sales_records "
                "(order_date, customer_name, product_code, product_name, quantity, unit_price, amount) "
                "VALUES (?, ?, ?, ?, ?, ?, ?)",
                (row["order_date"], row["customer_name"], row["product_code"],
                 row.get("product_name", ""), row["quantity"],
                 row.get("unit_price", 0.0), row.get("amount", 0.0)),
            )
        conn.commit()


def _minimal_sales_rows() -> list[dict]:
    return [
        {"order_date": f"2026-{m:02d}-{d:02d}", "customer_name": "A客戶",
         "product_code": "P1", "product_name": "商品A",
         "quantity": 20, "unit_price": 100.0, "amount": 2000.0}
        for m, d in [(3, 10), (3, 25), (4, 12), (4, 28)]
    ]


def test_normal_get_routes_do_not_call_pandas_read_excel(tmp_path, monkeypatch):
    """After DB-first migration, GET routes must not call pd.read_excel even if Excel exists."""
    # Create a dummy Excel file so file_path.exists() would be True
    config = ForecastConfig()
    dummy_excel = tmp_path / config.detail_file
    pd.DataFrame(columns=list(config.required_columns)).to_excel(dummy_excel, index=False, sheet_name=config.detail_sheet)

    excel_reads: list = []

    def guard_read_excel(*args, **kwargs):
        excel_reads.append(str(args[0]) if args else str(kwargs))
        raise AssertionError(f"pd.read_excel called during page request: {args[0] if args else kwargs}")

    monkeypatch.setattr(pd, "read_excel", guard_read_excel)

    db_base = _isolated_db_base()
    _seed_db_sales(db_base, _minimal_sales_rows())
    client = _client({"DB_BASE_PATH": db_base, "DATA_BASE_PATH": str(tmp_path)})

    routes = ["/", "/forecast", "/monitor/products", "/settings"]
    for route in routes:
        resp = client.get(route)
        assert resp.status_code == 200, f"{route} returned {resp.status_code}"
        assert not excel_reads, f"pd.read_excel called when serving {route}: {excel_reads}"


def test_export_route_does_not_call_pandas_read_excel(tmp_path, monkeypatch):
    """POST /export must not read Excel — it should derive the target from the DB."""
    config = ForecastConfig()
    dummy_excel = tmp_path / config.detail_file
    pd.DataFrame(columns=list(config.required_columns)).to_excel(dummy_excel, index=False, sheet_name=config.detail_sheet)

    excel_reads: list = []

    def guard_read_excel(*args, **kwargs):
        excel_reads.append(str(args[0]) if args else str(kwargs))
        raise AssertionError(f"pd.read_excel called during export: {args[0] if args else kwargs}")

    monkeypatch.setattr(pd, "read_excel", guard_read_excel)

    db_base = _isolated_db_base()
    _seed_db_sales(db_base, _minimal_sales_rows())
    client = _client({"DB_BASE_PATH": db_base, "DATA_BASE_PATH": str(tmp_path)})

    resp = client.post("/export", data={"year": "2026", "month": "5"})
    assert resp.status_code == 200
    assert not excel_reads, f"pd.read_excel called during export: {excel_reads}"
