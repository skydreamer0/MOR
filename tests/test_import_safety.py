"""Synthetic import contract regression tests; no company workbook data."""
from io import BytesIO
import pandas as pd
import pytest
from src.backend.etl import normalize_budget_targets
from src.backend.daily_sales_importer import normalize_daily_sales, DAILY_SALES_COLUMNS


def test_budget_preserves_encoded_year():
    rows = normalize_budget_targets(pd.DataFrame([{'客戶簡稱':'Synthetic','商品號':'0001','標  題':'07預算S總量',202601:2}]), None)
    assert (rows.iloc[0]['year'],rows.iloc[0]['month']) == (2026,1)


def test_numeric_daily_date_is_calendar_date():
    row={v:'' for v in DAILY_SALES_COLUMNS.values()}
    row.update({'出貨日期':20260130,'客戶簡稱':'Synthetic','產品':'0001', '銷售數量':2,'贈品數量':0,'銷貨淨價':1,'含稅淨額':2,'折後業績':2})
    assert str(normalize_daily_sales(pd.DataFrame([row])).sales_date.iloc[0]) == '2026-01-30'


def test_budget_unknown_metrics_fails():
    with pytest.raises(ValueError):
        normalize_budget_targets(pd.DataFrame([{'客戶簡稱':'Synthetic','商品號':'0001','年':2026,'標  題':'01銷+贈量',1:2}]),None)

from src.backend.import_parser import parse_workbook, strict_date, parse_period
from src.backend.import_service import preflight, commit_plan
from src.backend.database import MORDatabase
from src.backend.history_service import fetch_last_month_actuals, fetch_trend_6m
from src.backend.monthly_actuals import fetch_monthly_actuals
from src.backend.monthly_review_data import MonthlyReviewDataReader
from src.backend.data_loader import load_sales_detail_from_db
from src.backend.row_identity import make_row_id


def budget(customer='A',value=20,period=202601,metric='07預算S總量'):
    return pd.DataFrame([{'客戶簡稱':customer,'商品號':'0001','標  題':metric,period:value}])


def detail(customer='A',qty=2,day=1):
    return pd.DataFrame([{'年':2026,'月':1,'日':day,'客戶簡稱':customer,'商品號':'0001',
                          '商品簡稱':'Synthetic','銷+贈S量':qty,'單價NT(淨)':3,'含稅總額(淨)':qty*3}])


def monthly(customer='A',year=2025):
    return pd.DataFrame([{'年':year,'客戶簡稱':customer,'商品號':'0001','標  題':label,1:value}
                         for label,value in [('01銷+贈量',2),('02銷+贈S量',20),('03銷含稅總額',100),('04銷未稅總額',95),('05折後業績額',90)]])


def workbook(*frames,offset=0,right=0,notes=True):
    out=BytesIO()
    with pd.ExcelWriter(out,engine='openpyxl') as writer:
        if notes:pd.DataFrame({'備註':['Synthetic only']}).to_excel(writer,sheet_name='notes',index=False)
        for index,frame in enumerate(frames):
            frame.to_excel(writer,sheet_name=f'renamed {index}',index=False,startrow=offset,startcol=right)
    return out.getvalue()


def db_dump(db):
    with db.get_connection() as conn:return '\n'.join(conn.iterdump())


@pytest.mark.parametrize('value',[0,13,202613,1.5,'1.5',202601.5])
def test_invalid_periods_fail(value):
    with pytest.raises(ValueError):normalize_budget_targets(budget(period=value),2026)


@pytest.mark.parametrize('value',[None,'oops',float('inf'),float('nan'),'=A1'])
def test_invalid_monthly_values_fail(value):
    with pytest.raises(ValueError):normalize_budget_targets(budget(value=value),None)


@pytest.mark.parametrize('value',[20260230,45292,'2026/01/01',20260101.5])
def test_invalid_dates_fail(value):
    with pytest.raises(ValueError):strict_date(value)


def test_period_conflicts_and_missing_year_fail():
    with pytest.raises(ValueError):normalize_budget_targets(budget(),2027)
    with pytest.raises(ValueError):normalize_budget_targets(budget(period=1),None)
    frame=budget();frame['年']=2027
    with pytest.raises(ValueError):normalize_budget_targets(frame,None)
    frame=budget();frame['202601']=3
    with pytest.raises(ValueError):normalize_budget_targets(frame,None)


def test_layout_and_file_count_are_provenance_only():
    first=preflight([('a.xlsx',workbook(budget('A'),budget('B'),offset=4,right=2))])
    second=preflight([('renamed.xlsx',workbook(budget('B').iloc[:,::-1])),('other.xlsx',workbook(budget('A')))])
    assert first.package_hash==second.package_hash
    assert first.manifest[0]['header_row']==5
    assert first.manifest[0]['header_column']==3
    assert first.manifest[0]['source_columns']['商品號']=='D'


def test_same_sheet_separate_blocks_and_combined_types():
    out=BytesIO()
    with pd.ExcelWriter(out,engine='openpyxl') as writer:
        budget('A').to_excel(writer,index=False,startrow=3)
        detail('B').to_excel(writer,index=False,startrow=8)
    plan=preflight([('combined.xlsx',out.getvalue())])
    assert set(plan.frames)=={'budget','sales'}


def test_sample_budget_ancillary_labels_do_not_confuse_history():
    frame=pd.concat([budget(metric=k) for k in ['01銷+贈量','02預算總量','03銷售總額','04預算總額','05金達成率','06銷+贈S量','07預算S總量']])
    plan=preflight([('budget.xlsx',workbook(frame))])
    assert set(plan.frames)=={'budget'}
    assert len(plan.frames['budget'])==3


def test_duplicate_artifact_and_equal_monthly_policy():
    payload=workbook(budget())
    assert preflight([('a.xlsx',payload),('renamed.xlsx',payload)]).duplicate_artifacts==1
    other=workbook(budget(),offset=4)
    with pytest.raises(ValueError,match='Overlapping monthly'):preflight([('a',payload),('b',other)])
    assert len(preflight([('a',payload),('b',other)],deduplicate_monthly=True).frames['budget'])==1
    with pytest.raises(ValueError):preflight([('a',payload),('b',workbook(budget(value=99)))],deduplicate_monthly=True)


def test_daily_equal_transactions_are_preserved_and_partitions_explicit():
    frame=pd.concat([detail(),detail()],ignore_index=True)
    assert len(preflight([('a',workbook(frame))]).frames['sales'])==2
    with pytest.raises(ValueError,match='Overlapping daily'):preflight([('a',workbook(detail())),('b',workbook(detail(day=2)))])
    assert len(preflight([('a',workbook(detail())),('b',workbook(detail(day=2)))],daily_partitions=True).frames['sales'])==2


def test_no_input_unknown_only_and_uncached_formula_fail():
    with pytest.raises(ValueError):preflight([])
    with pytest.raises(ValueError):preflight([('wrong',workbook(pd.DataFrame({'unknown':[1]})))])
    with pytest.raises(ValueError,match='cached'):preflight([('formula',workbook(budget(value='=1+1')))])


def test_whole_package_preflight_has_no_side_effects(tmp_path):
    db=MORDatabase(tmp_path/'db');before=db_dump(db)
    with pytest.raises(ValueError):
        plan=preflight([('good',workbook(detail())),('bad',workbook(budget(value='damaged')))])
        commit_plan(db,plan)
    assert db_dump(db)==before


def test_atomic_cross_table_rollback(tmp_path,monkeypatch):
    import src.backend.import_service as service
    db=MORDatabase(tmp_path/'db');before=db_dump(db)
    plan=preflight([('both',workbook(detail(),budget()))])
    def fail(*args):raise RuntimeError('injected second-table failure')
    monkeypatch.setattr(service,'_write_budget',fail)
    with pytest.raises(RuntimeError):commit_plan(db,plan)
    assert db_dump(db)==before


def test_failure_in_receipt_rolls_back_facts(tmp_path):
    db=MORDatabase(tmp_path/'db')
    with db.get_connection() as conn:
        conn.execute("CREATE TRIGGER reject_receipt BEFORE INSERT ON import_receipts BEGIN SELECT RAISE(ABORT,'synthetic failure'); END")
    before=db_dump(db)
    with pytest.raises(Exception,match='synthetic failure'):commit_plan(db,preflight([('a',workbook(detail(),budget()))]))
    assert db_dump(db)==before


def test_narrow_budget_updates_preserve_other_months_customers_metrics(tmp_path):
    db=MORDatabase(tmp_path/'db')
    commit_plan(db,preflight([('a',workbook(budget('A'),budget('B'),budget('A',period=202602),budget('A',metric='04預算總額',value=100)))]))
    commit_plan(db,preflight([('new',workbook(budget('A',value=30)))]))
    with db.get_connection() as conn:
        rows=conn.execute('SELECT customer_name,month,target_quantity,target_amount FROM budget_targets ORDER BY customer_name,month').fetchall()
    assert [tuple(r) for r in rows]==[('A',1,30,100),('A',2,20,None),('B',1,20,None)]


def test_replay_idempotent_after_disjoint_scope_append(tmp_path):
    db=MORDatabase(tmp_path/'db')
    plan=preflight([('a',workbook(detail()))]);commit_plan(db,plan)
    commit_plan(db,preflight([('b',workbook(detail('B')))]))
    before=db_dump(db)
    assert commit_plan(db,plan)['changed'] is False
    assert db_dump(db)==before


def test_partial_overlap_fails_and_explicit_complete_scope_replaces(tmp_path):
    db=MORDatabase(tmp_path/'db')
    commit_plan(db,preflight([('a',workbook(detail(),detail('B')))]))
    before=db_dump(db)
    with pytest.raises(ValueError,match='complete_months'):commit_plan(db,preflight([('new',workbook(detail(qty=8)))]))
    assert db_dump(db)==before
    commit_plan(db,preflight([('new',workbook(detail(qty=8)))],complete_months={(2026,1)}))
    with db.get_connection() as conn:assert conn.execute('SELECT SUM(quantity),COUNT(*) FROM sales_records').fetchone()[:]==(8,1)


def test_closed_month_guard_checked_at_commit(tmp_path):
    db=MORDatabase(tmp_path/'db');plan=preflight([('a',workbook(budget(),detail()))])
    with db.get_connection() as conn:conn.execute('INSERT INTO month_close_records(year,month) VALUES (2026,1)')
    before=db_dump(db)
    with pytest.raises(ValueError,match='已結月'):commit_plan(db,plan)
    assert db_dump(db)==before


def test_monthly_requires_observed_periods_and_complete_metrics():
    with pytest.raises(ValueError,match='observed_months'):preflight([('a',workbook(monthly()))])
    with pytest.raises(ValueError,match='required metrics'):preflight([('a',workbook(monthly().iloc[:1]))],observed_months={(2025,1)})


def test_monthly_history_consumed_without_fabricated_orders_or_prices(tmp_path):
    db=MORDatabase(tmp_path/'db')
    commit_plan(db,preflight([('a',workbook(monthly()))],observed_months={(2025,1)}))
    key=make_row_id('A','0001')
    facts=fetch_monthly_actuals(db,2025,1)[key]
    assert facts['quantity']==20 and facts['base_quantity']==2
    assert facts['unit_price'] is None and facts['latest_order_day'] is None
    assert load_sales_detail_from_db(db).empty
    assert fetch_last_month_actuals(db,2025,2)=={key:20}
    assert fetch_trend_6m(db,2025,2)[key][-1]==20
    reader=MonthlyReviewDataReader(db)
    assert reader.last_year_rows(2026,1)[key]=={'qty':20,'amount':100}
    assert reader.customer_amounts(2025,1)=={'A':100}
    assert reader.product_totals(2025,1)[:2]==({'0001':100},{'0001':20})
    assert reader.fallback_unit_prices(2026,1)=={}


def test_monthly_daily_overlap_fails_in_both_directions(tmp_path):
    db=MORDatabase(tmp_path/'db')
    hist=monthly(year=2026)
    commit_plan(db,preflight([('a',workbook(hist))],observed_months={(2026,1)}))
    before=db_dump(db)
    with pytest.raises(ValueError,match='overlap'):commit_plan(db,preflight([('b',workbook(detail()))]))
    assert db_dump(db)==before
    with pytest.raises(ValueError,match='overlap'):preflight([('both',workbook(hist,detail()))],observed_months={(2026,1)})


def test_partial_schema_in_second_sheet_blocks_package():
    with pytest.raises(ValueError,match='incomplete'):
        preflight([('a',workbook(budget(),detail().drop(columns=['含稅總額(淨)'])))])


def test_blank_separator_inside_table_does_not_truncate_facts():
    frame=pd.concat([detail(),pd.DataFrame([{c:None for c in detail().columns}]),detail('B')],ignore_index=True)
    assert len(preflight([('a',workbook(frame))]).frames['sales'])==2


def test_formatted_numeric_identity_requires_text_instead_of_losing_zeroes():
    import openpyxl
    book=openpyxl.load_workbook(BytesIO(workbook(budget())))
    sheet=book['renamed 0'];sheet['B2']=123;sheet['B2'].number_format='00000'
    stream=BytesIO();book.save(stream)
    with pytest.raises(ValueError,match='leading zeros'):preflight([('a',stream.getvalue())])


def test_preview_explains_monthly_coverage_and_is_nonmutating(tmp_path):
    from src.backend.etl import sync_excel_to_db
    db=MORDatabase(tmp_path/'db');before=db_dump(db)
    result=sync_excel_to_db(db,tmp_path,files=[('a',workbook(monthly()))],preview=True)
    assert result['can_commit'] is False and 'observed_months' in result['unresolved'][0]
    assert result['tables'][0]['periods']==[(2025,1)]
    assert db_dump(db)==before


def test_commit_cannot_use_an_unresolved_preview(tmp_path):
    plan=preflight([('a',workbook(monthly()))],allow_unresolved=True)
    db=MORDatabase(tmp_path/'db')
    with pytest.raises(ValueError,match='observed_months'):commit_plan(db,plan)


def test_historical_promotion_rollback_keeps_current_rows(tmp_path):
    from src.backend.etl import import_current_month
    from src.backend.etl import SHPB_COLUMNS
    db=MORDatabase(tmp_path/'db')
    row={'出貨日期':20260101,'客戶簡稱':'A','產品':'0001','產品簡稱':'Synthetic','銷售數量':2,'贈品數量':0,'銷貨淨價':3,'含稅淨額':6}
    import_current_month(db,pd.DataFrame([row]))
    with db.get_connection() as conn:
        conn.execute("CREATE TRIGGER fail_retirement BEFORE DELETE ON current_month_records BEGIN SELECT RAISE(ABORT,'retirement failure'); END")
    before=db_dump(db)
    with pytest.raises(Exception,match='retirement failure'):
        commit_plan(db,preflight([('a',workbook(detail()))],complete_months={(2026,1)}))
    assert db_dump(db)==before


def test_legacy_db_overlap_is_reported_instead_of_double_counted(tmp_path):
    db=MORDatabase(tmp_path/'db')
    with db.get_connection() as conn:
        for table in ['sales_records','current_month_records']:
            conn.execute(f"INSERT INTO {table}(order_date,customer_name,product_code,quantity) VALUES ('2026-01-01','A','0001',2)")
    with pytest.raises(ValueError,match='overlap'):load_sales_detail_from_db(db)


def test_conflicting_customer_name_identity_blocks_package():
    a=budget();a['客戶代號']='C1'
    b=budget('A',period=202602);b['客戶代號']='C2'
    with pytest.raises(ValueError,match='customer code/name'):
        preflight([('a',workbook(a,b))])


def test_failure_during_executemany_rolls_back_prior_rows(tmp_path):
    db=MORDatabase(tmp_path/'db')
    with db.get_connection() as conn:
        conn.execute("CREATE TRIGGER fail_second BEFORE INSERT ON sales_records WHEN NEW.quantity=99 BEGIN SELECT RAISE(ABORT,'last row failure'); END")
    before=db_dump(db)
    plan=preflight([('a',workbook(pd.concat([detail(),detail('B',qty=99)])))])
    with pytest.raises(Exception,match='last row failure'):commit_plan(db,plan)
    assert db_dump(db)==before


def test_concurrent_close_is_seen_after_write_lock_acquisition(tmp_path):
    import sqlite3
    from concurrent.futures import ThreadPoolExecutor
    from threading import Event
    db=MORDatabase(tmp_path/'db');plan=preflight([('a',workbook(budget()))])
    closing=sqlite3.connect(db.db_path)
    closing.execute('BEGIN IMMEDIATE')
    closing.execute('INSERT INTO month_close_records(year,month) VALUES (2026,1)')
    started=Event()
    def importing():
        started.set()
        return commit_plan(db,plan)
    with ThreadPoolExecutor(max_workers=1) as pool:
        future=pool.submit(importing)
        assert started.wait(2)
        closing.commit();closing.close()
        with pytest.raises(ValueError,match='已結月'):future.result(timeout=5)
    with db.get_connection() as conn:
        assert conn.execute('SELECT COUNT(*) FROM budget_targets').fetchone()[0]==0


def test_provenance_retains_hash_header_and_fact_coordinates(tmp_path):
    import json
    db=MORDatabase(tmp_path/'db')
    plan=preflight([('private-name.xlsx',workbook(budget(),offset=4,right=2))]);commit_plan(db,plan)
    with db.get_connection() as conn:
        manifest=json.loads(conn.execute('SELECT manifest FROM import_receipts').fetchone()[0])
        sources=json.loads(conn.execute('SELECT sources FROM import_fact_sources').fetchone()[0])
    assert len(manifest[0]['sha256'])==64
    assert manifest[0]['header_row']==5
    assert 'F6' in sources['cells'][0]


def test_cell_error_and_duplicate_headers_reject_with_coordinates():
    import openpyxl
    book=openpyxl.load_workbook(BytesIO(workbook(budget())))
    sheet=book['renamed 0'];sheet['D2']='#DIV/0!';sheet['D2'].data_type='e'
    out=BytesIO();book.save(out)
    with pytest.raises(ValueError,match='D2'):parse_workbook(out.getvalue(),'broken.xlsx')
    book=openpyxl.load_workbook(BytesIO(workbook(budget())))
    sheet=book['renamed 0'];sheet['E1']='客戶簡稱';sheet['E2']='A'
    out=BytesIO();book.save(out)
    with pytest.raises(ValueError,match='duplicate headers'):parse_workbook(out.getvalue(),'duplicate.xlsx')


def test_workbook_resource_bound_is_explicit(monkeypatch):
    import src.backend.import_parser as parser
    monkeypatch.setattr(parser,'MAX_CELLS',2)
    with pytest.raises(ValueError,match='scanned-cell limit'):parse_workbook(workbook(budget()),'large.xlsx')


def test_ambiguous_legacy_product_alias_is_rejected_instead_of_rekeyed(tmp_path):
    db=MORDatabase(tmp_path/'db');frame=detail();frame['商品號']='001.0'
    with pytest.raises(ValueError,match='canonical product code'):
        commit_plan(db,preflight([('a',workbook(frame))]))
    assert load_sales_detail_from_db(db).empty


def test_explicit_monthly_deduplication_also_applies_within_table():
    frame=pd.concat([budget(),budget()],ignore_index=True)
    with pytest.raises(ValueError,match='Overlapping monthly'):preflight([('a',workbook(frame))])
    assert len(preflight([('a',workbook(frame))],deduplicate_monthly=True).frames['budget'])==1


def test_understated_sheet_dimensions_cannot_hide_facts():
    import zipfile
    import re
    payload=workbook(pd.concat([detail(),detail('B')],ignore_index=True),notes=False)
    modified=BytesIO()
    with zipfile.ZipFile(BytesIO(payload)) as source,zipfile.ZipFile(modified,'w') as target:
        for item in source.infolist():
            data=source.read(item.filename)
            if item.filename=='xl/worksheets/sheet1.xml':
                data=re.sub(rb'<dimension ref="[^"]+"',b'<dimension ref="A1:A1"',data)
            target.writestr(item,data)
    assert len(preflight([('a',modified.getvalue())]).frames['sales'])==2


def test_cross_package_customer_identity_conflict_is_atomic(tmp_path):
    db=MORDatabase(tmp_path/'test.db')
    first=budget();first['客戶代號']='C1'
    commit_plan(db,preflight([('first.xlsx',workbook(first))]))
    before=db_dump(db)
    second=budget(value=99);second['客戶代號']='C2'
    with pytest.raises(ValueError,match='identity|mapping'):
        commit_plan(db,preflight([('second.xlsx',workbook(second))]))
    assert db_dump(db)==before
    second['客戶代號']='C1'
    assert commit_plan(db,preflight([('second.xlsx',workbook(second))]))['changed']


@pytest.mark.parametrize('header',['-1','+1',' 1','.5','\t202601'])
def test_signed_or_padded_period_never_silently_skipped(header):
    frame=budget();frame[header]=999
    with pytest.raises(ValueError):normalize_budget_targets(frame,2026)


def test_daily_provenance_survives_append_replay_and_replacement(tmp_path):
    import json
    db=MORDatabase(tmp_path/'test.db')
    def upload(name,customer,**kwargs):
        return commit_plan(db,preflight([(name,workbook(detail(customer)))],**kwargs))
    a=upload('first.xlsx','A');b=upload('second.xlsx','B')
    def sources():
        with db.get_connection() as conn:
            return {tuple(json.loads(r['fact_key'])):r['receipt_id'] for r in conn.execute("SELECT * FROM import_fact_sources WHERE kind='sales'")}
    assert sources()=={(2026,1,'A','0001'):a['receipt_id'],(2026,1,'B','0001'):b['receipt_id']}
    assert not upload('replay.xlsx','A')['changed']
    assert len(sources())==2
    replacement=upload('replacement.xlsx','C',complete_months={(2026,1)})
    assert sources()=={(2026,1,'C','0001'):replacement['receipt_id']}


def test_monthly_history_enriches_primary_comparison_arrays(tmp_path):
    from src.backend.forecast_engine import build_forecast,ForecastTarget,ForecastOptions
    from src.backend.history_service import enrich_rows_with_history
    db=MORDatabase(tmp_path/'test.db')
    commit_plan(db,preflight([('history.xlsx',workbook(monthly(),monthly('MonthlyOnly')))],observed_months={(2025,1)}))
    commit_plan(db,preflight([('daily.xlsx',workbook(detail()))]))
    summary=build_forecast(load_sales_detail_from_db(db),ForecastTarget(2026,1),ForecastOptions(include_all=True))
    row=enrich_rows_with_history(summary.rows,db,2026,1)[0]
    assert row.ly_monthly[0]==20 and row.ly_monthly_amount[0]==100
    assert row.last_year_same_month_qty==20
    assert row.latest_price==3 and row.latest_order_date.isoformat()=='2026-01-01'
    # Monthly-only entities remain available at true monthly grain, not synthetic daily rows.
    fact=fetch_monthly_actuals(db,2025,1)['MonthlyOnly__0001']
    assert fact['quantity']==20 and fact['unit_price'] is None and fact['latest_order_day'] is None


def test_unverified_legacy_customer_binding_requires_resolution(tmp_path):
    db=MORDatabase(tmp_path/'test.db')
    commit_plan(db,preflight([('legacy.xlsx',workbook(budget()))]))
    before=db_dump(db)
    coded=budget(period=202602);coded['客戶代號']='C1'
    with pytest.raises(ValueError,match='unverified'):
        commit_plan(db,preflight([('coded.xlsx',workbook(coded))]))
    assert db_dump(db)==before


def test_customer_binding_survives_disjoint_periods_and_rolls_back(tmp_path):
    db=MORDatabase(tmp_path/'test.db')
    first=budget();first['客戶代號']='C1'
    commit_plan(db,preflight([('first.xlsx',workbook(first))]))
    other=budget(period=202602);other['客戶代號']='C1'
    assert commit_plan(db,preflight([('other.xlsx',workbook(other))]))['changed']
    other['客戶代號']='C2'
    with pytest.raises(ValueError,match='identity'):
        commit_plan(db,preflight([('conflict.xlsx',workbook(other))]))
    new=budget(customer='B');new['客戶代號']='C3'
    with db.get_connection() as conn:
        conn.execute("CREATE TRIGGER reject_receipt BEFORE INSERT ON import_receipts BEGIN SELECT RAISE(ABORT,'synthetic failure'); END")
    before=db_dump(db)
    with pytest.raises(Exception,match='synthetic failure'):
        commit_plan(db,preflight([('failure.xlsx',workbook(new))]))
    assert db_dump(db)==before


def test_monthly_comparisons_reach_workbench_context(tmp_path):
    from src.backend.forecast_workbench_context import build
    from src.backend.forecast_config import ForecastConfig
    db=MORDatabase(tmp_path/'test.db')
    prior=monthly();current=monthly(year=2026).rename(columns={1:2})
    commit_plan(db,preflight([('history.xlsx',workbook(prior,current))],observed_months={(2025,1),(2026,2)}))
    commit_plan(db,preflight([('daily.xlsx',workbook(detail()))]))
    context=build(ForecastConfig(),db,{'year':'2026','month':'1'})
    row=context.rows[0]
    assert row.ly_monthly[0]==20 and row.ly_monthly_amount[0]==100
    assert row.ty_monthly[1]==20 and row.ty_monthly_amount[1]==100
    assert row.last_year_same_month_qty==20
